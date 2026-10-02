"""Testes do pacote EF sem IA real: o próprio gerador monta uma EF no template do
projeto, e o leitor/validador precisam reconhecê-la (o ciclo gerar → validar fecha)."""
import os
from pathlib import Path

import pytest

import opuscore_ef as ef
from opuscore_core.sdk.testes import IAFalsaSync
from opuscore_ef.geracao.gerar_ef import RascunhoEF

RAIZ = Path(__file__).resolve().parents[3]
TEMPLATE = RAIZ / "config" / "templates" / "EF_template.docx"

RASC = {"identificacao": {"id_gap": "MM-201", "descricao_gap": "Bloqueio de pedido sem certificação", "modulo": "MM",
                          "cenario": "Compras", "processo": "Pedido de compra"},
        "resumo": {"tipos_programa": ["Enhancement"], "prioridade": "Alta", "alternativa_standard": "[A CONFIRMAR]",
                   "inventario": [{"objeto": "ZCL_MM_CERT_CHECK", "tipo": "Enhancement", "natureza": "NOVO", "objetivo": "Bloquear"}]},
        "objetivo": {"as_is": "Planilha", "problema": "Pedido sem certificação", "to_be": "Bloquear gravação",
                     "beneficio": "Conformidade", "criterio_sucesso": "[A CONFIRMAR]"},
        "processos": [{"processo": "Criar pedido", "transacao": "ME21N", "app_fiori": "", "papel": "Comprador"}],
        "regras": [{"id": "RN-01", "condicao": "Fornecedor sem certificação", "acao": "Impedir gravação", "excecao": "Grupo de exceção"}],
        "escopo": {"incluido": ["ME21N"], "fora_escopo": ["ZTMM_FORA_ESCOPO"], "premissas": ["Tabela nova"], "dependencias": []},
        "fluxo": [{"passo": 1, "ator": "Comprador", "sistema": "S/4HANA", "acao": "Cria o pedido"}],
        "sistemas": [{"sistema": "S/4HANA", "tipo": "SAP", "papel": "Pedido", "integracao": "Não"}],
        "testes": [{"cenario": "Sem certificação", "passos": "ME21N", "resultado_esperado": "Bloqueio"}],
        "complementares": {"volumetria": "[A CONFIRMAR]"}, "pontos_a_confirmar": ["ME22N?"]}


@pytest.fixture(scope="module")
def rascunho(tmp_path_factory):
    os.environ["OPUSCORE_DATA"] = str(tmp_path_factory.mktemp("dados"))
    ia = IAFalsaSync([RASC])
    r, caminho, uso = ef.gerar_rascunho("Workshop: bloquear pedido...", nome_workshop="WS.docx", template=TEMPLATE,
                                        destino=tmp_path_factory.mktemp("out") / "MM-201.docx", provider=ia, id_gap="MM-201")
    return r, caminho


def test_rascunho_herda_cabecalho_e_destaca_a_confirmar(rascunho):
    import docx
    from docx.enum.text import WD_COLOR_INDEX
    r, caminho = rascunho
    d = docx.Document(str(caminho))
    assert sum(1 for x in d.sections[-1].header.part.rels.values() if "image" in x.reltype) == 2
    amarelos = [run for t in d.tables for row in t.rows for c in row.cells for p in c.paragraphs for run in p.runs
                if run.font.highlight_color == WD_COLOR_INDEX.YELLOW]
    assert amarelos and isinstance(r, RascunhoEF)


def test_validador_reconhece_as_secoes_do_rascunho(rascunho):
    from opuscore_ef.validacao import validar as v
    _, caminho = rascunho
    res, para_ia = v.preparar(ef.ler(caminho), ef.regras_padrao())
    ids = {s["id"] for s in para_ia}
    assert {"capa", "resumo", "objetivo", "regras", "fluxo", "nfr", "teste_proc"} <= ids
    assert not [x for x in res if x.origem == "regra" and x.status == "CRITICO"]


def test_regras_em_camadas():
    regras = ef.regras_padrao({"regras": {"prompt": "do módulo"}}, {"regras": {"ativo": False}})
    r = next(x for x in regras if x["id"] == "regras")
    assert r["prompt"] == "do módulo" and r["ativo"] is False and r["editado"]
    assert len(regras) == 22


def test_validar_secoes_marca_trecho_inventado(rascunho):
    _, caminho = rascunho
    rev = {"secoes": [{"id": "regras", "status": "ATENCAO", "resumo": "ok",
                       "achados": [{"tipo": "AMBIGUIDADE", "descricao": "x", "trecho": "frase que não existe", "sugestao": "y"}]}]}
    res, uso = ef.validar_secoes(caminho, provider=IAFalsaSync([rev]), usar_cache=False)
    regra = next(x for x in res if x.id == "regras")
    assert regra.status == "ATENCAO" and regra.achados[0]["trecho_confirmado"] is False


def test_etapa1_sem_ia_e_objetos(rascunho):
    _, caminho = rascunho
    r = ef.analisar(caminho, estado="RASCUNHO", usar_ia=False)
    assert r.gate.status and "##" in r.feedback_markdown
    no_escopo, fora = ef.objetos_tecnicos(caminho)
    assert "ZCL_MM_CERT_CHECK" in no_escopo
    assert ef.refs_z_texto("usa ZFOO e ZBAR\n* ZCOMENTADO") == ["ZBAR", "ZFOO"]


@pytest.mark.skipif(not os.environ.get("EF_SD034"), reason="defina EF_SD034=<caminho da EF SD-034> para o teste de referência")
def test_referencia_sd034():
    no_escopo, fora = ef.objetos_tecnicos(Path(os.environ["EF_SD034"]))
    assert len(no_escopo) == 9 and fora == ["ZTMMD_TOPAGEM_P"]
