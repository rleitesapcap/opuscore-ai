"""Upload de documentos (.docx/.pdf) com validação dupla: extensão + assinatura real
do arquivo + texto extraível. Usado por qualquer consultor (EF, Workshop B...)."""
from __future__ import annotations

import io
import json
import os
import re
import uuid
import zipfile
from pathlib import Path

from fastapi import APIRouter, File, HTTPException, UploadFile

EF_ALLOWED_EXT = {".docx", ".pdf"}
EF_MAX_BYTES = 25 * 1024 * 1024  # 25 MB


class DocumentoInvalido(ValueError):
    """Documento recusado: formato, conteúdo ou tamanho inválido."""


def _store() -> Path:
    from opuscore_core.sdk.config import diretorio_dados
    return diretorio_dados() / "uploads"


def _docx_text(data: bytes) -> str:
    import docx  # python-docx
    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(dict.fromkeys(cells)))
    return "\n".join(parts)


def _pdf_text(data: bytes) -> str:
    # mesma ordem de leitores do motor de ET: pdfplumber -> pypdf
    try:
        import pdfplumber
        out = []
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                t = page.extract_text() or ""
                if t.strip():
                    out.append(t)
        if out:
            return "\n".join(out)
    except Exception:
        pass
    import pypdf
    reader = pypdf.PdfReader(io.BytesIO(data))
    return "\n".join(t for t in ((pg.extract_text() or "") for pg in reader.pages) if t.strip())


def validar(filename: str, data: bytes) -> dict:
    """Valida o documento. Levanta DocumentoInvalido com mensagem clara se recusada.

    Regras: extensão .docx ou .pdf; assinatura real do arquivo compatível com a
    extensão (um .txt renomeado é recusado); tamanho <= 25 MB; texto extraível
    (PDF escaneado sem camada de texto é recusado, pois o motor não conseguiria
    ler a EF). Devolve metadados + prévia do texto.
    """
    name = os.path.basename(filename or "").strip()
    ext = Path(name).suffix.lower()
    if ext == ".doc":
        raise DocumentoInvalido("Formato .doc (Word 97-2003) não é aceito. Salve como .docx ou exporte para PDF.")
    if ext not in EF_ALLOWED_EXT:
        raise DocumentoInvalido(f"Formato não aceito ({ext or 'sem extensão'}). A EF precisa ser .docx ou .pdf.")
    if not data:
        raise DocumentoInvalido("Arquivo vazio.")
    if len(data) > EF_MAX_BYTES:
        raise DocumentoInvalido(f"Arquivo acima do limite de {EF_MAX_BYTES // (1024 * 1024)} MB.")

    if ext == ".pdf":
        if not data[:1024].lstrip().startswith(b"%PDF-"):
            raise DocumentoInvalido("O arquivo tem extensão .pdf mas não é um PDF válido.")
        try:
            text = _pdf_text(data)
        except ImportError:
            raise DocumentoInvalido("Leitor de PDF não instalado no servidor (pip install pdfplumber pypdf).")
        except Exception as e:  # noqa: BLE001
            raise DocumentoInvalido(f"Não foi possível ler o PDF: {e}")
        if not text.strip():
            raise DocumentoInvalido(
                "O PDF não tem texto extraível (provavelmente é escaneado/imagem). "
                "Envie a EF em .docx ou um PDF gerado a partir do documento.")
    else:
        if not data.startswith(b"PK"):
            raise DocumentoInvalido("O arquivo tem extensão .docx mas não é um documento Word válido.")
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if "word/document.xml" not in z.namelist():
                    raise DocumentoInvalido("O arquivo .docx não contém um documento Word (word/document.xml ausente).")
        except zipfile.BadZipFile:
            raise DocumentoInvalido("O arquivo .docx está corrompido.")
        try:
            text = _docx_text(data)
        except ImportError:
            raise DocumentoInvalido("Leitor de Word não instalado no servidor (pip install python-docx).")
        except Exception as e:  # noqa: BLE001
            raise DocumentoInvalido(f"Não foi possível ler o .docx: {e}")
        if not text.strip():
            raise DocumentoInvalido("O .docx não tem texto. Verifique se é a EF correta.")

    return {
        "filename": name, "ext": ext, "size": len(data), "chars": len(text),
        "preview": text.strip()[:600], "text": text,
    }


def _safe_name(name: str) -> str:
    base = re.sub(r"[^A-Za-z0-9._ -]+", "_", os.path.basename(name)).strip() or "EF"
    return base[:150]


def guardar(filename: str, data: bytes, meta: dict) -> str:
    """Grava o documento validado em <dados>/uploads/<id>/ preservando o nome original."""
    ef_id = uuid.uuid4().hex
    folder = _store() / ef_id
    folder.mkdir(parents=True, exist_ok=True)
    fname = _safe_name(filename)
    (folder / fname).write_bytes(data)
    info = {k: v for k, v in meta.items() if k != "text"}
    info["stored_name"] = fname
    (folder / "meta.json").write_text(json.dumps(info, ensure_ascii=False), encoding="utf-8")
    (folder / "ef_text.txt").write_text(meta.get("text", ""), encoding="utf-8")
    return ef_id


def carregar(ef_id: str) -> tuple[Path, dict, str]:
    """(caminho do arquivo original, metadados, texto extraído) de um documento armazenado."""
    if not re.fullmatch(r"[0-9a-f]{32}", ef_id or ""):
        raise DocumentoInvalido("Identificador de documento inválido.")
    folder = _store() / ef_id
    meta_p = folder / "meta.json"
    if not meta_p.exists():
        raise DocumentoInvalido("Documento não encontrado. Faça o upload novamente.")
    meta = json.loads(meta_p.read_text(encoding="utf-8"))
    path = folder / meta["stored_name"]
    text = (folder / "ef_text.txt").read_text(encoding="utf-8") if (folder / "ef_text.txt").exists() else ""
    return path, meta, text


class ServicoUploads:
    """Implementação de ctx.uploads."""

    def carregar(self, upload_id: str) -> tuple[Path, dict, str]:
        from opuscore_core.contratos.erros import invalido
        try:
            return carregar(upload_id)
        except DocumentoInvalido as e:
            raise invalido(str(e))


router = APIRouter(prefix="/api/plataforma", tags=["uploads"])


@router.post("/uploads")
async def upload(file: UploadFile = File(...)):
    data = await file.read(EF_MAX_BYTES + 1)
    try:
        meta = validar(file.filename or "", data)
    except DocumentoInvalido as e:
        raise HTTPException(400, str(e))
    uid = guardar(file.filename or "documento", data, meta)
    return {"upload_id": uid, "ef_id": uid, "filename": meta["filename"], "ext": meta["ext"],
            "size": meta["size"], "chars": meta["chars"], "preview": meta["preview"]}
