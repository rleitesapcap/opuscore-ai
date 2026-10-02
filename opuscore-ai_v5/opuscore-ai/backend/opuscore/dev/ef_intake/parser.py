"""Segmenta a EF (.docx/.pdf) em seções com localização rastreável.

DOCX: a hierarquia vem dos ESTILOS de título (Heading 1/2/3), não da numeração
manual, que costuma estar inconsistente. Cada bloco recebe uma localização:
  P087          -> parágrafo 87 do corpo do documento
  T2.r11        -> tabela 2, linha 11
PDF: não há títulos confiáveis; cada página vira uma seção e cada linha um bloco:
  PG3.L12       -> página 3, linha 12
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class Block:
    loc: str
    kind: str          # "p" | "table" | "line"
    text: str


@dataclass
class Section:
    sid: str           # S00, S01...
    level: int         # 0 = conteúdo antes do primeiro título
    title: str
    path: list[str]    # títulos ancestrais + o próprio
    blocks: list[Block] = field(default_factory=list)
    images: int = 0

    @property
    def path_str(self) -> str:
        return " > ".join(self.path)

    def text(self) -> str:
        return "\n".join(b.text for b in self.blocks)


@dataclass
class EFDocument:
    filename: str
    fmt: str                                   # "docx" | "pdf"
    sections: list[Section]
    header: dict = field(default_factory=dict)  # rótulo -> valor (tabela de capa)
    revisions: list[dict] = field(default_factory=list)
    total_images: int = 0

    # ---- consultas usadas pelas verificações ----
    def locs(self) -> dict[str, tuple[Section, Block]]:
        return {b.loc: (s, b) for s in self.sections for b in s.blocks}

    def section_by_id(self, sid: str) -> Section | None:
        return next((s for s in self.sections if s.sid == sid), None)

    def find_loc_containing(self, needle: str) -> str | None:
        n = _norm(needle)
        for s in self.sections:
            for b in s.blocks:
                if n and n in _norm(b.text):
                    return b.loc
        return None

    def to_prompt_text(self) -> str:
        out = []
        for s in self.sections:
            img = f" — contém {s.images} imagem(ns) não legível(is) em texto" if s.images else ""
            out.append(f"\n## [{s.sid}] {s.path_str}{img}")
            for b in s.blocks:
                out.append(f"[{b.loc}] {b.text}")
        return "\n".join(out).strip()


def _norm(s: str) -> str:
    s = (s or "").replace("\u00a0", " ").replace("“", '"').replace("”", '"').replace("’", "'")
    return re.sub(r"\s+", " ", s).strip().lower()


_HEADING = re.compile(r"(?i)^(heading|t[ií]tulo)\s*(\d)")
_NUM_PREFIX = re.compile(r"^\s*\d+(\.\d+)*\.?\s*")


def _clean_title(t: str) -> str:
    return re.sub(r"\s+", " ", _NUM_PREFIX.sub("", t)).strip()


def _row_cells(row) -> list[str]:
    cells: list[str] = []
    for c in row.cells:
        x = re.sub(r"\s+", " ", c.text.replace("\n", " / ")).strip()
        if x and (not cells or cells[-1] != x):   # células mescladas repetem o texto
            cells.append(x)
    return cells


def _images_in(element) -> int:
    xml = element.xml if hasattr(element, "xml") else ""
    return xml.count("<pic:pic") + xml.count("<v:imagedata")


def parse_docx(path: str | Path) -> EFDocument:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(str(path))
    sections: list[Section] = [Section("S00", 0, "Capa e cabeçalho", ["Capa e cabeçalho"])]
    stack: list[tuple[int, str]] = []
    pi = ti = 0
    total_images = 0
    header: dict = {}
    revisions: list[dict] = []

    for ch in d.element.body.iterchildren():
        tag = ch.tag.split("}")[1]
        if tag == "p":
            p = Paragraph(ch, d)
            txt = re.sub(r"\s+", " ", p.text).strip()
            imgs = _images_in(ch)
            total_images += imgs
            m = _HEADING.match(p.style.name or "")
            if m and txt:
                level = int(m.group(2))
                title = _clean_title(txt)
                stack = [x for x in stack if x[0] < level] + [(level, title)]
                sections.append(Section(f"S{len(sections):02d}", level, title, [t for _, t in stack]))
            elif txt:
                sections[-1].blocks.append(Block(f"P{pi:03d}", "p", txt))
            sections[-1].images += imgs
            pi += 1
        elif tag == "tbl":
            tb = Table(ch, d)
            total_images += _images_in(ch)
            sections[-1].images += _images_in(ch)
            for r, row in enumerate(tb.rows):
                cells = _row_cells(row)
                if not cells:
                    continue
                line = " | ".join(cells)
                sections[-1].blocks.append(Block(f"T{ti}.r{r}", "table", line))
                # cabeçalho "Rótulo: valor"
                for c in cells:
                    mm = re.match(r"^([^:]{2,40}):\s*(.+)$", c)
                    if mm and sections[-1].sid == "S00":
                        header.setdefault(mm.group(1).strip(), mm.group(2).strip())
                # histórico de revisão: "V0 | autor | data | descrição"
                if re.match(r"^V\d+$", cells[0]) and len(cells) >= 3:
                    revisions.append({"versao": cells[0], "autor": cells[1],
                                      "data": cells[2] if len(cells) > 2 else "",
                                      "descricao": cells[3] if len(cells) > 3 else "",
                                      "loc": f"T{ti}.r{r}"})
            ti += 1

    sections = [s for s in sections if s.blocks or s.images or s.level > 0]
    return EFDocument(Path(path).name, "docx", sections, header, revisions, total_images)


def parse_pdf(path: str | Path) -> EFDocument:
    sections: list[Section] = []
    try:
        import pdfplumber
        with pdfplumber.open(str(path)) as pdf:
            for n, page in enumerate(pdf.pages, start=1):
                s = Section(f"S{n:02d}", 1, f"Página {n}", [f"Página {n}"], images=len(page.images))
                for i, line in enumerate((page.extract_text() or "").splitlines(), start=1):
                    if line.strip():
                        s.blocks.append(Block(f"PG{n}.L{i}", "line", line.strip()))
                sections.append(s)
    except ImportError:
        import pypdf
        reader = pypdf.PdfReader(str(path))
        for n, page in enumerate(reader.pages, start=1):
            s = Section(f"S{n:02d}", 1, f"Página {n}", [f"Página {n}"])
            for i, line in enumerate((page.extract_text() or "").splitlines(), start=1):
                if line.strip():
                    s.blocks.append(Block(f"PG{n}.L{i}", "line", line.strip()))
            sections.append(s)
    header: dict = {}
    for s in sections[:2]:
        for b in s.blocks:
            mm = re.match(r"^([^:]{2,40}):\s*(.+)$", b.text)
            if mm:
                header.setdefault(mm.group(1).strip(), mm.group(2).strip())
    return EFDocument(Path(path).name, "pdf", sections, header, [],
                      sum(s.images for s in sections))


def parse_ef(path: str | Path) -> EFDocument:
    ext = Path(path).suffix.lower()
    if ext == ".docx":
        return parse_docx(path)
    if ext == ".pdf":
        return parse_pdf(path)
    raise ValueError(f"Formato de EF não suportado: {ext} (use .docx ou .pdf)")
