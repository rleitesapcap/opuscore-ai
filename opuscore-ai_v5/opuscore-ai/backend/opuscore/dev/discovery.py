"""Descoberta e orquestração dos fluxos SAP-sourced do Desenvolvedor.

A origem do código deixa de ser um diretório e passa a ser o S/4 (ADT, via MCP):
- Remediação: a partir da EF, descobre os objetos + o crawl de dependências Z,
  o usuário aprova a lista, e só então o source é buscado no S/4 e remediado.
- ET: objetos vêm de uma REQUEST de transporte OU de uma lista informada; a EF é
  OBRIGATÓRIA e chega como ARQUIVO (.docx ou .pdf, validado); o source é buscado
  no S/4 e a ET é gerada.

Os motores incorporados (dev/remediation e dev/et) NÃO são alterados: nós apenas
montamos, num diretório temporário, exatamente o layout que cada motor espera e
apontamos o pipeline para lá.

Layout esperado pelo motor de ET (extract.py):
    <root>/01_EF/        -> a EF original (.docx/.pdf), lida pelos leitores do motor
    <root>/02_CODIGOS/   -> um arquivo por objeto, com sufixo ADT (ex.: .clas.abap)
Os nomes das pastas respeitam DIR_EF / DIR_CODIGOS se definidos no ambiente.
"""
from __future__ import annotations

import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import uuid
import zipfile
from pathlib import Path

from ..sap.adt_client import _extract_z_refs

DEV_DIR = Path(__file__).resolve().parent
ET_DIR = DEV_DIR / "et"
REM_DIR = DEV_DIR / "remediation"
BACKEND_DIR = DEV_DIR.parents[1]
EF_STORE = BACKEND_DIR / "data" / "uploads" / "ef"

EF_ALLOWED_EXT = {".docx", ".pdf"}
EF_MAX_BYTES = 25 * 1024 * 1024  # 25 MB


class EFValidationError(ValueError):
    """EF recusada: formato, conteúdo ou tamanho inválido."""


# ---------------------------------------------------------------------------
# EF: validação (extensão + assinatura real + texto extraível) e armazenamento
# ---------------------------------------------------------------------------
def _docx_text(data: bytes) -> str:
    import docx  # python-docx
    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(dict.fromkeys(cells)))
    return "\n".join(parts)


def _pdf_text(data: bytes) -> str:
    # mesma ordem de leitores do motor de ET: pdfplumber -> pypdf
    try:
        import pdfplumber
        out = []
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                t = page.extract_text() or ""
                if t.strip():
                    out.append(t)
        if out:
            return "\n".join(out)
    except Exception:
        pass
    import pypdf
    reader = pypdf.PdfReader(io.BytesIO(data))
    return "\n".join(t for t in ((pg.extract_text() or "") for pg in reader.pages) if t.strip())


def validate_ef(filename: str, data: bytes) -> dict:
    """Valida a EF. Levanta EFValidationError com mensagem clara se recusada.

    Regras: extensão .docx ou .pdf; assinatura real do arquivo compatível com a
    extensão (um .txt renomeado é recusado); tamanho <= 25 MB; texto extraível
    (PDF escaneado sem camada de texto é recusado, pois o motor não conseguiria
    ler a EF). Devolve metadados + prévia do texto.
    """
    name = os.path.basename(filename or "").strip()
    ext = Path(name).suffix.lower()
    if ext == ".doc":
        raise EFValidationError("Formato .doc (Word 97-2003) não é aceito. Salve como .docx ou exporte para PDF.")
    if ext not in EF_ALLOWED_EXT:
        raise EFValidationError(f"Formato não aceito ({ext or 'sem extensão'}). A EF precisa ser .docx ou .pdf.")
    if not data:
        raise EFValidationError("Arquivo vazio.")
    if len(data) > EF_MAX_BYTES:
        raise EFValidationError(f"Arquivo acima do limite de {EF_MAX_BYTES // (1024 * 1024)} MB.")

    if ext == ".pdf":
        if not data[:1024].lstrip().startswith(b"%PDF-"):
            raise EFValidationError("O arquivo tem extensão .pdf mas não é um PDF válido.")
        try:
            text = _pdf_text(data)
        except ImportError:
            raise EFValidationError("Leitor de PDF não instalado no servidor (pip install pdfplumber pypdf).")
        except Exception as e:  # noqa: BLE001
            raise EFValidationError(f"Não foi possível ler o PDF: {e}")
        if not text.strip():
            raise EFValidationError(
                "O PDF não tem texto extraível (provavelmente é escaneado/imagem). "
                "Envie a EF em .docx ou um PDF gerado a partir do documento.")
    else:
        if not data.startswith(b"PK"):
            raise EFValidationError("O arquivo tem extensão .docx mas não é um documento Word válido.")
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if "word/document.xml" not in z.namelist():
                    raise EFValidationError("O arquivo .docx não contém um documento Word (word/document.xml ausente).")
        except zipfile.BadZipFile:
            raise EFValidationError("O arquivo .docx está corrompido.")
        try:
            text = _docx_text(data)
        except ImportError:
            raise EFValidationError("Leitor de Word não instalado no servidor (pip install python-docx).")
        except Exception as e:  # noqa: BLE001
            raise EFValidationError(f"Não foi possível ler o .docx: {e}")
        if not text.strip():
            raise EFValidationError("O .docx não tem texto. Verifique se é a EF correta.")

    return {
        "filename": name, "ext": ext, "size": len(data), "chars": len(text),
        "preview": text.strip()[:600], "text": text,
    }


def _safe_name(name: str) -> str:
    base = re.sub(r"[^A-Za-z0-9._ -]+", "_", os.path.basename(name)).strip() or "EF"
    return base[:150]


def store_ef(filename: str, data: bytes, meta: dict) -> str:
    """Grava a EF validada em data/uploads/ef/<ef_id>/ preservando o nome original."""
    ef_id = uuid.uuid4().hex
    folder = EF_STORE / ef_id
    folder.mkdir(parents=True, exist_ok=True)
    fname = _safe_name(filename)
    (folder / fname).write_bytes(data)
    info = {k: v for k, v in meta.items() if k != "text"}
    info["stored_name"] = fname
    (folder / "meta.json").write_text(json.dumps(info, ensure_ascii=False), encoding="utf-8")
    (folder / "ef_text.txt").write_text(meta.get("text", ""), encoding="utf-8")
    return ef_id


def load_ef(ef_id: str) -> tuple[Path, dict, str]:
    """(caminho do arquivo original, metadados, texto extraído) de uma EF armazenada."""
    if not re.fullmatch(r"[0-9a-f]{32}", ef_id or ""):
        raise EFValidationError("Identificador de EF inválido.")
    folder = EF_STORE / ef_id
    meta_p = folder / "meta.json"
    if not meta_p.exists():
        raise EFValidationError("EF não encontrada. Faça o upload novamente.")
    meta = json.loads(meta_p.read_text(encoding="utf-8"))
    path = folder / meta["stored_name"]
    text = (folder / "ef_text.txt").read_text(encoding="utf-8") if (folder / "ef_text.txt").exists() else ""
    return path, meta, text


# ---------------------------------------------------------------------------
# Checagem prévia da IA dos motores (mesmas variáveis do llm_providers.py)
# ---------------------------------------------------------------------------
def llm_preflight(motor: str = "et") -> tuple[bool, str]:
    """Confere se as variáveis que o motor exige estão definidas. Não chama a IA.

    ATENÇÃO: o motor de remediação (lc_CapRemediation.py) usa um cliente próprio
    que fala SOMENTE com a Generative Engine da Capgemini, ignorando LLM_PROVIDER.
    Por isso, para a remediação, as credenciais da Capgemini são sempre exigidas."""
    if motor == "remediation":
        faltam = [v for v in ("LLM_API_KEY", "WORKSPACE_ID") if not os.environ.get(v, "").strip()]
        if faltam:
            return False, ("O motor de remediação usa somente a Generative Engine da Capgemini "
                           f"(independe do LLM_PROVIDER) e exige {', '.join(faltam)} no backend/.env.")
        return True, "capgemini"
    prov = (os.environ.get("LLM_PROVIDER") or "capgemini").strip().lower()
    if prov in ("capgemini", "generative-engine", "ge"):
        faltam = [v for v in ("LLM_API_KEY", "WORKSPACE_ID") if not os.environ.get(v, "").strip()]
    elif prov in ("anthropic", "anthropic-api", "api"):
        faltam = [] if os.environ.get("ANTHROPIC_API_KEY", "").strip() else ["ANTHROPIC_API_KEY"]
    elif prov in ("claude_code", "claude-code", "cli", "cc"):
        faltam = [] if shutil.which("claude") else ["CLI 'claude' no PATH"]
    else:
        return False, f"LLM_PROVIDER '{prov}' desconhecido. Use capgemini | anthropic | claude_code."
    if faltam:
        return False, (f"IA dos motores não configurada: LLM_PROVIDER={prov} exige "
                       f"{', '.join(faltam)} no backend/.env.")
    return True, prov


# ---------------------------------------------------------------------------
# Objetos e materialização de source
# ---------------------------------------------------------------------------
def extract_ef_objects(ef_text: str) -> list[str]:
    """Identifica os objetos (Z*/Y*) citados num TEXTO de EF (uso via API com texto colado)."""
    return _extract_z_refs(ef_text or "")


def objetos_da_ef(ef_path: Path) -> tuple[list[str], list[str]]:
    """Objetos custom (Z*/Y*) da EF para a remediação, usando o extrator determinístico
    da Etapa 1: nomes quebrados em tabela são reparados, campos (TABELA-CAMPO) não viram
    objeto e o que a EF declara fora do escopo é separado.
    Retorna (objetos_no_escopo, objetos_fora_do_escopo)."""
    from .ef_intake import deterministic as det
    from .ef_intake.parser import parse_ef

    itens, _ = det.extrair(parse_ef(ef_path))
    no_escopo, fora = [], []
    for i in itens:
        v = (i.valor_original or "").upper()
        if i.categoria != "REF_TECNICA" or not re.match(r"^[ZY]", v) or "-" in v or \
                i.tipo_objeto in ("CAMPO_TABELA", "METODO"):
            continue
        (fora if i.escopo == "EXCLUIDO" else no_escopo).append(v)
    return sorted(dict.fromkeys(no_escopo)), sorted(dict.fromkeys(fora))


# sufixo por tipo ADT — o motor de ET identifica o objeto pelo sufixo composto
_SUFFIX_BY_TYPE = {
    "CLAS": ".clas.abap", "INTF": ".intf.abap", "DDLS": ".ddls.asddls",
    "BDEF": ".bdef.asbdef", "DDLX": ".ddlx.asddlx", "SRVB": ".srvb.asrvb",
    "SRVD": ".asrvd", "DCLS": ".dcl",
}


def _fname(name: str, adt_type: str = "") -> str:
    prefix = (adt_type or "").split("/")[0].upper()
    return name.strip() + _SUFFIX_BY_TYPE.get(prefix, ".abap")


def write_sources(items: list[dict], dest: Path) -> list[str]:
    """items = [{'name','source','type'?}]. Grava um arquivo por objeto."""
    dest.mkdir(parents=True, exist_ok=True)
    written = []
    for it in items:
        src = it.get("source") or ""
        if not src.strip():
            continue
        (dest / _fname(it["name"], it.get("type", ""))).write_text(src, encoding="utf-8")
        written.append(it["name"])
    return written


def ler_uso(caminho: Path) -> list[dict]:
    """Linhas gravadas pelos motores em LLM_USAGE_LOG (uma por chamada de IA)."""
    usos = []
    if caminho.is_file():
        for linha in caminho.read_text(encoding="utf-8").splitlines():
            try:
                usos.append(json.loads(linha))
            except json.JSONDecodeError:
                pass
    return usos


def resumo_uso(caminho: Path) -> dict:
    from ..usage import resumo
    usos = ler_uso(caminho)
    modelo = next((u.get("modelo") for u in usos if u.get("modelo")), None)
    return resumo(usos, modelo)


def run_pipeline(script: Path, args: list[str], cwd: Path, env_extra: dict) -> tuple[int, str]:
    """Executa um pipeline incorporado como subprocesso (isola estado, herda o
    .env do backend via ambiente). Retorna (rc, cauda da saída combinada)."""
    env = os.environ.copy()
    env.update(env_extra or {})
    # Os motores leem o modelo com nomes diferentes: a Etapa 2 da ET (generate_docs.py)
    # usa LLM_MODEL; a montagem final e o .env do projeto usam LLM_MODEL_MAIN. Sem
    # LLM_MODEL, o motor cai no modelo fixo do código, que pode não existir na conta.
    if not (env.get("LLM_MODEL") or "").strip() and (env.get("LLM_MODEL_MAIN") or "").strip():
        env["LLM_MODEL"] = env["LLM_MODEL_MAIN"].strip()
    proc = subprocess.run(
        [sys.executable, str(script), *args],
        cwd=str(cwd), env=env, capture_output=True, text=True, timeout=1800,
    )
    full = ((proc.stdout or "") + "\n" + (proc.stderr or "")).strip()
    return proc.returncode, full[-60000:]


_DIAG = re.compile(r"(\[✗ FAIL\].*|Abortando:.*|FALHOU.*|Traceback.*|\w*Error: .*|PIPELINE INTERROMPIDO.*)")


def diagnose(log: str) -> list[str]:
    """Linhas que explicam uma falha do motor (validação da IA, etapa que falhou,
    exceções), na ordem em que aparecem — para mostrar em destaque na tela."""
    seen, out = set(), []
    for line in (log or "").splitlines():
        m = _DIAG.search(line)
        if m:
            txt = re.sub(r"^\d{2}:\d{2}:\d{2} \[\w+\]\s*", "", m.group(1)).strip()
            if txt and txt not in seen:
                seen.add(txt); out.append(txt[:400])
    return out[:8]


def collect(folder: Path, patterns: tuple[str, ...]) -> list[Path]:
    out: list[Path] = []
    for pat in patterns:
        out += sorted(folder.rglob(pat), key=lambda p: p.stat().st_mtime, reverse=True)
    return out


# ---------------------------------------------------------------------------
# Execução dos motores (origem = diretório materializado do S/4)
# ---------------------------------------------------------------------------
def run_remediation(sources: list[dict]) -> tuple[int, str, list[Path], dict]:
    """sources = [{'name','source'}] já aprovados. Materializa e roda o motor."""
    work = Path(tempfile.mkdtemp(prefix="opuscore_rem_"))
    in_dir, out_dir = work / "input", work / "output"
    write_sources(sources, in_dir)
    uso_log = work / "uso_ia.jsonl"
    rc, log = run_pipeline(
        REM_DIR / "run_pipeline.py", [], REM_DIR,
        {"REMEDIATION_INPUT_DIR": str(in_dir), "REMEDIATION_OUTPUT_DIR": str(out_dir),
         "LLM_USAGE_LOG": str(uso_log)},
    )
    return rc, log, collect(out_dir, ("*.abap", "*.txt", "*.md")), resumo_uso(uso_log)


def prepare_et_workspace(sources: list[dict], ef_path: Path) -> tuple[Path, Path]:
    """Monta o layout do motor de ET: <root>/01_EF (EF original) e <root>/02_CODIGOS."""
    work = Path(tempfile.mkdtemp(prefix="opuscore_et_"))
    root, out_dir = work / "input", work / "output"
    ef_dir = root / os.environ.get("DIR_EF", "01_EF")
    codes_dir = root / os.environ.get("DIR_CODIGOS", "02_CODIGOS")
    ef_dir.mkdir(parents=True, exist_ok=True)
    shutil.copy2(ef_path, ef_dir / ef_path.name)   # arquivo ORIGINAL, sem conversão
    write_sources(sources, codes_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    return root, out_dir


def run_et(sources: list[dict], ef_path: Path) -> tuple[int, str, list[Path], dict]:
    """sources = [{'name','source','type'}] + EF (arquivo .docx/.pdf). Roda o motor da ET."""
    root, out_dir = prepare_et_workspace(sources, ef_path)
    uso_log = root.parent / "uso_ia.jsonl"
    # o pipeline confere a ET em ET_OUTBOUND_DIR/DIR_ET_GERADA; a Etapa 3 precisa gravar lá
    et_gerada = out_dir / os.environ.get("DIR_ET_GERADA", "01_ET_GERADA")
    et_gerada.mkdir(parents=True, exist_ok=True)
    rc, log = run_pipeline(
        ET_DIR / "run_et_pipeline.py",
        ["--root", str(root), "--output", str(et_gerada)], ET_DIR,
        {
            "ET_OUTBOUND_DIR": str(out_dir),
            "DIR_ET_GERADA": et_gerada.name,
            # fixa as pastas: o .env.example do motor usa outros nomes
            # (Functional-Specification/Codigos); a plataforma monta 01_EF/02_CODIGOS
            "DIR_EF": os.environ.get("DIR_EF", "01_EF"),
            "DIR_CODIGOS": os.environ.get("DIR_CODIGOS", "02_CODIGOS"),
            "DIR_TEMPLATE_LOCAL": "rag",   # template Word oficial fica em dev/et/rag/
            "LLM_USAGE_LOG": str(uso_log),
        },
    )
    return rc, log, collect(out_dir, ("ET_*.docx", "ET_*.md", "*.docx")), resumo_uso(uso_log)
