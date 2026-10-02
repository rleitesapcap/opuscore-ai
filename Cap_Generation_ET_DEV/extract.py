#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
===============================================================================
 extract.py — O EXTRATOR MECÂNICO (a primeira etapa, sem Inteligência Artificial)
===============================================================================

O QUE ESTE ARQUIVO FAZ:
    É o ponto de partida da geração de documentação técnica SAP. Antes de
    qualquer IA entrar em ação, este script varre os arquivos de entrada (a EF —
    Especificação Funcional — e os códigos ABAP) e monta, de forma determinística
    (ou seja: sempre o mesmo resultado para a mesma entrada, sem "criatividade"),
    o material bruto que a IA vai usar depois: um "esqueleto" de documento para
    cada objeto, uma lista de requisitos extraídos da EF, e um arquivo
    `_payload.json` com tudo organizado.

POR QUE FAZER ISSO SEM IA:
    Tudo que pode ser descoberto por regras simples (tipo do objeto, blocos de
    remediação, campos de tela de seleção, variáveis TVARV, objetos de
    autorização, etc.) é extraído aqui em Python puro, usando leitura de arquivo
    e regex (um jeito de "procurar padrões de texto", explicado abaixo onde
    aparece pela primeira vez). Isso é mais rápido, mais barato e mais confiável
    do que pedir para a IA "ler" tudo isso — a IA entra só depois, para escrever
    a PROSA (as explicações em linguagem natural) em cima do que já foi extraído.

RAIZ DO PROJETO:
    É sempre a pasta onde está este arquivo (onde também fica o .env).
    input é a pasta de ENTRADA (inbound): a EF e os códigos entram direto
    nela, e o script NUNCA escreve dentro dela — só lê.

  INBOUND  {projeto}/input/01_EF
           {projeto}/input/02_CODIGOS
  OUTBOUND {projeto}/{ET_OUTBOUND_DIR}/{DIR_ET_GERADA}/{DIR_ARTEFATOS}/

Uso:
  python extract.py
  python extract.py --date DDMMYYYY
===============================================================================
"""

import argparse    # lê os parâmetros passados na linha de comando (ex.: --date)
import json        # lê/grava dados no formato JSON (o _payload.json de saída)
import os          # funções de sistema operacional: caminhos, pastas, arquivos
import re          # regex — "expressões regulares", um jeito de buscar padrões de texto
import sys         # acesso a stderr e sys.exit(), para erros e código de saída
from dataclasses import dataclass  # dataclass: cria uma classe só para guardar dados, com menos código
from datetime import datetime      # data/hora, usado para nomear e carimbar os arquivos gerados


# ----------------------------------------------------------------------------
# .env — inbound (entrada) vs outbound (saída). Sempre relativo à raiz do projeto.
# ----------------------------------------------------------------------------

# Nomes padrão das pastas/arquivos do fluxo, usados quando o .env não define um
# valor customizado para a chave (veja a função env_name logo abaixo).
_DIR_DEFAULTS = {
    "ET_INBOUND_DIR": "input",
    "DIR_EF": "01_EF",
    "DIR_TEMP_ET": "02_TEMP_ET",
    "DIR_CODIGOS": "02_CODIGOS",
    "DIR_CODE_REVIEW": "03_CODE_REVIEW",
    "ET_OUTBOUND_DIR": "output",
    "DIR_ET_GERADA": "01_ET_GERADA",
    "DIR_ARTEFATOS": "artefatos",
    "FILE_REGRAS": "regras_extrair.md",
    "FILE_REQUISITOS": "requisitos_extraidos.md",
    "FILE_ACHADOS": "achados_tecnicos.md",
    "FILE_PENDENCIAS": "pendencias_review.md",
    "FILE_RESULTADO": "resultado_final.md",
    "DIR_TEMPLATE_LOCAL": "rap",
    "ET_TEMPLATE_NAME": "",
}


def load_dotenv(path):
    """Lê um arquivo .env (linhas 'CHAVE=valor') e coloca cada chave nas
    variáveis de ambiente do processo, sem sobrescrever o que já existir."""
    if not path or not os.path.isfile(path):
        return
    with open(path, encoding="utf-8") as fh:
        for line in fh:
            line = line.strip()
            # ignora linhas vazias, comentários (#) e linhas sem "=" (inválidas)
            if not line or line.startswith("#") or "=" not in line:
                continue
            k, v = line.split("=", 1)
            # setdefault: só define a variável se ela AINDA não existir no ambiente
            # (assim, uma variável já exportada no shell tem prioridade sobre o .env)
            os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))


def project_dir():
    """Raiz do projeto: pasta do extract.py. Nunca input."""
    return os.path.dirname(os.path.abspath(__file__))


def load_project_env(root=None):  # noqa: ARG001 — compat; nunca lê .env do inbound
    """Lê só {projeto}/.env. Ignora qualquer .env dentro de input."""
    load_dotenv(os.path.join(project_dir(), ".env"))


def env_name(key):
    """Devolve o valor customizado no .env para essa chave, ou o padrão em
    _DIR_DEFAULTS se ela não tiver sido definida (ou estiver vazia)."""
    v = os.environ.get(key)
    if v and str(v).strip():
        return str(v).strip()
    return _DIR_DEFAULTS[key]


@dataclass
class EtLayout:
    """Caminhos do lote. inbound = input (leitura); outbound = output (escrita).

    @dataclass gera automaticamente o "construtor" (__init__) a partir dos
    campos abaixo, então basta declarar os nomes e os tipos."""
    project: str        # raiz do projeto (pasta do extract.py)
    inbound: str        # pasta de entrada (normalmente {projeto}/input)
    outbound: str       # pasta de saída (normalmente {projeto}/output)
    ef: str              # pasta onde fica a EF (Especificação Funcional)
    codes: str           # pasta onde ficam os códigos ABAP de entrada
    review: str          # pasta com os artefatos do Code Inspector
    artifacts: str       # pasta final onde os documentos gerados são gravados
    artifacts_rel: str   # a mesma pasta 'artifacts', mas em caminho relativo (para exibir em log)

    # @property transforma um método em algo que se lê como se fosse um campo
    # comum (obj.ef_name, sem parênteses), mas o valor é calculado na hora.
    @property
    def ef_name(self):
        """Só o NOME da subpasta da EF (ex.: '01_EF'), sem o caminho completo."""
        return env_name("DIR_EF")

    @property
    def codes_name(self):
        """Só o NOME da subpasta dos códigos (ex.: '02_CODIGOS')."""
        return env_name("DIR_CODIGOS")

    @property
    def review_name(self):
        """Só o NOME da subpasta de review (ex.: '03_CODE_REVIEW')."""
        return env_name("DIR_CODE_REVIEW")


class LayoutError(ValueError):
    """Erro específico para quando a estrutura de pastas do projeto está errada
    (ex.: apontar --root para a própria pasta do extract.py). Herdar de
    ValueError significa que também pode ser tratado como um erro de valor comum."""
    pass


def _is_inside(path, parent):
    """True se `path` é `parent` ou está dentro dele."""
    path = os.path.normpath(os.path.abspath(path))
    parent = os.path.normpath(os.path.abspath(parent))
    try:
        return os.path.commonpath([path, parent]) == parent
    except ValueError:
        return False


def parse_location(root=None, project=None):
    """Monta o layout. Inbound = input (sem pasta extra).

    Sem flags   → lê {projeto}/input/{01_EF, 02_CODIGOS, ...}
    --root DIR  → lê DIR como inbound (não pode ser a pasta do extract.py)
    """
    project = os.path.abspath(project or project_dir())
    inbound_name = env_name("ET_INBOUND_DIR")
    outbound_name = env_name("ET_OUTBOUND_DIR")
    et_name = env_name("DIR_ET_GERADA")
    art_name = env_name("DIR_ARTEFATOS")
    default_inbound = os.path.join(project, inbound_name)
    outbound = os.path.join(project, outbound_name)

    if root:
        # --root foi passado: usa essa pasta como inbound, mas primeiro confere
        # que não é, por engano, a própria pasta do projeto (proteção comum).
        root = os.path.abspath(root)
        if root == project or os.path.isfile(os.path.join(root, "extract.py")):
            raise LayoutError(
                "a raiz do projeto é a pasta do extract.py. "
                "Coloque a EF em input/01_EF e os códigos em input/02_CODIGOS. "
                "Rode: python extract.py"
            )
        inbound = root
    else:
        inbound = default_inbound

    if os.path.normpath(outbound) == os.path.normpath(default_inbound):
        raise LayoutError(
            "ET_OUTBOUND_DIR não pode ser igual a ET_INBOUND_DIR. "
            "input é só entrada; a saída é  output/."
        )

    artifacts_rel = f"{outbound_name}/{et_name}/{art_name}"
    artifacts = os.path.join(outbound, et_name, art_name)
    # Rede de segurança: nunca deixa a pasta de saída (artifacts) ficar DENTRO
    # da pasta de entrada — isso poderia sobrescrever/misturar dados de origem.
    if _is_inside(artifacts, inbound) or _is_inside(artifacts, default_inbound):
        raise LayoutError(
            "recusando escrita dentro do inbound. "
            "Artefatos vão em ET_OUTBOUND_DIR (output/), não em input/."
        )
    return EtLayout(
        project=project,
        inbound=inbound,
        outbound=outbound,
        ef=os.path.join(inbound, env_name("DIR_EF")),
        codes=os.path.join(inbound, env_name("DIR_CODIGOS")),
        review=os.path.join(inbound, env_name("DIR_CODE_REVIEW")),
        artifacts=artifacts,
        artifacts_rel=artifacts_rel,
    )


def resolve_subdir(root, env_key, must_exist=False):
    """Subpasta inbound (DIR_EF, DIR_CODIGOS, ...) dentro de input."""
    name = env_name(env_key)
    return os.path.join(root, name), name


def resolve_inbound_root(base_dir=None):
    """Caminho completo da pasta de entrada (input) + o nome dela sozinho."""
    p = os.path.join(base_dir or project_dir(), env_name("ET_INBOUND_DIR"))
    return p, env_name("ET_INBOUND_DIR")


def resolve_outbound_root(base_dir=None):
    """Caminho completo da pasta de saída (output) + o nome dela sozinho."""
    p = os.path.join(base_dir or project_dir(), env_name("ET_OUTBOUND_DIR"))
    return p, env_name("ET_OUTBOUND_DIR") 


def resolve_et_gerada(create=False, base_dir=None):
    """{projeto}/output/01_ET_GERADA."""
    out, name = resolve_outbound_root(base_dir)
    p = os.path.join(out, env_name("DIR_ET_GERADA"))
    if create:
        os.makedirs(p, exist_ok=True)
    return p, f"{name}/{env_name('DIR_ET_GERADA')}"


def resolve_et_artifacts(create=False, base_dir=None, root=None):
    """Só OUTBOUND: {projeto}/output/01_ET_GERADA/artefatos/."""
    lay = parse_location(root=root, project=base_dir)
    if create:
        os.makedirs(lay.artifacts, exist_ok=True)
    return lay.artifacts, lay.artifacts_rel


def llm_enabled():
    """True se o pipeline vai usar a LLM (workspace Capgemini)."""
    if os.getenv("ET_DRY_RUN", "").strip().lower() in {"1", "true", "yes"}:
        return False
    provider = os.getenv("LLM_PROVIDER", "capgemini").strip().lower()
    if provider in {"off", "none", "local", "disabled"}:
        return False
    return bool(os.getenv("WORKSPACE_ID", "").strip())


def find_local_template(folder, name=""):
    """Primeiro .docx/.pdf de template na pasta local (rap/ na raiz)."""
    if not folder or not os.path.isdir(folder):
        return None
    files = []
    for fn in os.listdir(folder):
        if fn.startswith(".") or fn.startswith("~$"):
            continue
        ext = os.path.splitext(fn)[1].lower()
        if ext not in {".docx", ".doc", ".pdf"}:
            continue
        files.append(fn)
    if not files:
        return None
    if name:
        for fn in files:
            if fn == name or fn.lower() == name.lower():
                return os.path.join(folder, fn)

    def score(fn):
        """Pontua o quão provável é que este arquivo seja o template certo,
        pelo nome (quanto maior a pontuação, mais "parece" um template)."""
        low = fn.lower()
        s = 0
        if "template" in low:
            s += 10
        if "especifica" in low or "_et_" in low:
            s += 5
        if "tecnica" in low or "técnica" in low:
            s += 3
        return (-s, low)   # negativo para ordenar do MAIOR score para o menor
    files.sort(key=score)
    return os.path.join(folder, files[0])


def resolve_et_template(use_llm=None, base_dir=None):
    """Com LLM: template no WORKSPACE_ID. Sem LLM: arquivo em rap/ (raiz do projeto)."""
    base_dir = base_dir or project_dir()
    if use_llm is None:
        use_llm = llm_enabled()
    name = (os.environ.get("ET_TEMPLATE_NAME") or "").strip()
    if use_llm:
        # com LLM habilitada, o template mora no workspace remoto (na base de
        # conhecimento da Capgemini) — não precisamos procurar um arquivo local.
        return {
            "source": "workspace",
            "workspace_id": (os.environ.get("WORKSPACE_ID") or "").strip(),
            "name": name,
            "path": None,
            "dir": None,
        }
    used = env_name("DIR_TEMPLATE_LOCAL") or "rap"
    folder = os.path.join(base_dir, used)
    path = None
    if name:
        cand = os.path.join(folder, name)
        if os.path.isfile(cand):
            path = cand
    if not path:
        path = find_local_template(folder, name)
    return {
        "source": "local",
        "workspace_id": None,
        "name": os.path.basename(path) if path else (name or None),
        "path": path,
        "dir": used,
    }


def describe_et_template(tpl):
    """Uma linha para log: workspace WORKSPACE_ID / nome, ou rap/arquivo."""
    tpl = tpl or {}
    if tpl.get("source") == "workspace":
        return (
            f"workspace {tpl.get('workspace_id') or '(WORKSPACE_ID)'} "
            f"/ {tpl.get('name') or '(ET_TEMPLATE_NAME)'}"
        )
    if tpl.get("path"):
        return f"local {tpl.get('dir')}/{tpl.get('name')}"
    folder = tpl.get("dir") or env_name("DIR_TEMPLATE_LOCAL")
    return (
        f"local {folder}/ "
        "(arquivo não encontrado; coloque o .docx em rap/ na raiz do projeto)"
    )


# A ORDEM DE DESENVOLVIMENTO: cada tipo de objeto SAP tem um "rank" (número).
# Documentos são listados/gerados do menor rank para o maior — ou seja, na
# ordem em que costuma-se implementar as camadas SAP (tabela antes de CDS,
# CDS antes de comportamento, comportamento antes de serviço, etc.).
TYPE_ORDER = [
    ("Tabela (DDL)",                 10),
    ("Estrutura/Data Element (DDIC)",12),
    ("CDS Interface View",           20),
    ("CDS Composition View",         25),
    ("CDS Consumption View",         30),
    ("CDS Projection View",          35),
    ("CDS View (genérico)",          36),
    ("Metadata Extension",           40),
    ("Access Control (DCL)",         42),
    ("Behavior Definition (Interface)", 45),
    ("Behavior Implementation (ABP/Handler)", 50),
    ("Behavior Definition (Projection)", 55),
    ("Service Definition",           60),
    ("Service Binding",              65),
    ("AMDP",                         70),
    ("Classe (Global)",              80),
    ("Interface OO",                 82),
    ("Function Group",               85),
    ("Function Module",              88),
    ("Report/Programa",              90),
    ("Top Include",                  92),
    ("Screen Include (PBO/PAI)",     94),
    ("Form/Rotina Include",          96),
    ("Include (genérico)",           98),
    ("Enhancement/BAdI",             99),
    ("Objeto (não classificado)",    120),
]
# Dicionário {nome do tipo: rank}, montado a partir da lista acima — um
# "list comprehension"/dict comprehension: monta um dicionário novo percorrendo
# TYPE_ORDER em uma linha só, em vez de um loop for tradicional.
TYPE_ORDER_MAP = {name: rank for name, rank in TYPE_ORDER}


# Regex pré-compiladas (re.compile guarda o padrão já "pronto para uso", o que
# é mais rápido do que recompilar toda vez que a função roda) usadas para
# reconhecer trechos típicos de cada tipo de objeto SAP dentro do código-fonte.
_PROJECTION_STMT = re.compile(r"^\s*projection\s*;", re.IGNORECASE | re.MULTILINE)
# Reconhece uma implementação de Behavior Handler: herda de
# cl_abap_behavior_handler, OU é uma classe "lhc_..." (convenção SAP), OU tem
# "for behavior of" (associa a classe a uma Behavior Definition).
_BEHAVIOR_HANDLER = re.compile(
    r"inheriting\s+from\s+cl_abap_behavior_handler|\bclass\s+lhc_|\bfor\s+behavior\s+of\b",
    re.IGNORECASE,
)
# Reconhece uma CDS View "de Interface" (convenção de nome ZI_/I_ nas Views raiz).
_CDS_INTERFACE = re.compile(
    r"define\s+(?:root\s+)?view\s+entity\s+(?:/\w+/)?(?:zi_|i_)",
    re.IGNORECASE,
)


def classify(path, text):
    """Tipo do objeto SAP a partir do FONTE (e da extensão ADT), não da pasta.

    É uma cadeia de if/elif: cada condição procura uma "impressão digital" do
    tipo de objeto (extensão de arquivo, ou uma palavra-chave típica do ABAP/CDS
    no início de uma linha) e devolve assim que encontra a primeira que bate.
    A ordem importa — condições mais específicas vêm antes das mais genéricas."""
    lower = os.path.basename(path).lower()
    low = text.lower()
    if lower.endswith(".dcl") or re.search(r"^\s*define\s+role\b", low, re.M):
        return "Access Control (DCL)"
    if lower.endswith((".ddls.asddls", ".asddls", ".ddls")):
        return classify_cds(text)
    if lower.endswith((".bdef", ".asbdef", ".bdef.asbdef")):
        if _PROJECTION_STMT.search(low) or "as projection on" in low:
            return "Behavior Definition (Projection)"
        return "Behavior Definition (Interface)"
    if lower.endswith((".srvd", ".asrvd")) or "define service" in low:
        return "Service Definition"
    if lower.endswith((".srvb", ".asrvb")):
        return "Service Binding"
    if lower.endswith((".ddlx", ".asddlx")) or "annotate view" in low or "annotate entity" in low:
        return "Metadata Extension"
    if re.search(r"^\s*define\s+table\s+entity", low, re.M) or re.search(r"^\s*define\s+structure", low, re.M):
        return "Tabela (DDL)"
    if "define view" in low or "define root view" in low or "define abstract entity" in low:
        return classify_cds(text)
    if re.search(r"^\s*managed\b|^\s*unmanaged\b|^\s*projection\s*;", low, re.M) \
            or "implementation in class" in low:
        if _PROJECTION_STMT.search(low) or "as projection on" in low:
            return "Behavior Definition (Projection)"
        return "Behavior Definition (Interface)"
    # Filename hints before CLASS DEFINITION (local classes live in TOP/includes).
    if "_top" in lower or re.search(r"top\.(abap|txt|prog)$", lower):
        return "Top Include"
    if re.search(r"_(pbo|pai)\.(abap|txt|prog)$", lower):
        return "Screen Include (PBO/PAI)"
    if _BEHAVIOR_HANDLER.search(low):
        return "Behavior Implementation (ABP/Handler)"
    if ("by database procedure" in low and "amdp" in low) or "cl_amdp" in low:
        return "AMDP"
    if re.search(r"^\s*class\s+\w+\s+definition", low, re.M):
        return "Classe (Global)"
    if re.search(r"^\s*interface\s+\w+", low, re.M) and "define" not in low.split("interface", 1)[0][-40:]:
        if "endinterface" in low or "methods" in low:
            return "Interface OO"
    if re.search(r"^\s*function-pool\b", low, re.M) or re.search(r"^\s*function\s+group", low, re.M):
        return "Function Group"
    if re.search(r"^\s*function\s+\w+\.", low, re.M) or "endfunction" in low:
        return "Function Module"
    if re.search(r"^\s*report\b|^\s*program\b", low, re.M):
        return "Report/Programa"
    if re.search(r"module\s+status_\d+\s+output", low) or re.search(r"module\s+user_command_\d+\s+input", low) \
       or "process before output" in low or "process after input" in low:
        return "Screen Include (PBO/PAI)"
    if re.search(r"^\s*form\s+\w+", low, re.M):
        return "Form/Rotina Include"
    if re.search(r"^\s*enhancement\b|badi|get badi", low, re.M):
        return "Enhancement/BAdI"
    if lower.endswith((".abap", ".txt", ".src", ".prog", ".asinc", ".asprog")):
        return "Include (genérico)"
    return "Objeto (não classificado)"


def classify_cds(text):
    """Refina o tipo de uma CDS View: Projection, Consumption, Composition ou
    Interface, olhando anotações e palavras-chave típicas de cada subtipo."""
    low = text.lower()
    if "provider contract transactional_query" in low or "as projection on" in low:
        return "CDS Projection View"
    if "@consumption" in low or "@ui" in low or "@odata.publish" in low:
        return "CDS Consumption View"
    if "composition of" in low or "association [1..*]" in low:
        return "CDS Composition View"
    if _CDS_INTERFACE.search(low):
        return "CDS Interface View"
    return "CDS View (genérico)"


# Marcador que aparece no topo de um fonte já remediado (corrigido) pela
# esteira de IA da Capgemini — é assim que sabemos que um objeto é uma
# "Remediação" e não um objeto totalmente novo.
REMEDIATION_HEADER = re.compile(r"Capgemini\s+AI\s+Remediation\s+Code\s*-\s*Brazil", re.IGNORECASE)
# Os marcadores BEGIN/END que envolvem cada trecho alterado (ver apply_changes.py,
# que é quem os escreve). Aqui só os LEMOS, para localizar os blocos modificados.
BEGIN_MOD = re.compile(r'^\s*[*"]*\s*BEGIN\s+OF\s+MODIFICATION\s*-\s*Capgemini\s+SAP\s+AI\s+Remediation.*$', re.IGNORECASE)
END_MOD = re.compile(r'^\s*[*"]*\s*END\s+OF\s+MODIFICATION\s*-\s*Capgemini\s+SAP\s+AI\s+Remediation.*$', re.IGNORECASE)
CONTEXT_LINES = 3   # quantas linhas de "contexto" mostrar antes/depois de cada bloco modificado


def is_remediation(text):
    """True se o fonte já foi remediado antes (tem o cabeçalho da Capgemini nas
    primeiras 60 linhas). Só olha o início do arquivo, por ser mais rápido."""
    return bool(REMEDIATION_HEADER.search("\n".join(text.splitlines()[:60])))


def find_location(lines, idx):
    """Descobre 'onde' (em qual METHOD/FORM/MODULE/FUNCTION) está a linha idx,
    olhando de baixo para cima até achar a declaração que abre aquele bloco.
    Se não achar nenhuma, é porque a linha está fora de qualquer sub-rotina."""
    pat = re.compile(r'^\s*(METHOD|FORM|MODULE|FUNCTION|CLASS|ENDCLASS|PERFORM)\s+([\w/]+)', re.IGNORECASE)
    for i in range(idx, -1, -1):
        m = pat.match(lines[i])
        if m and m.group(1).upper() in ("METHOD", "FORM", "MODULE", "FUNCTION"):
            return f"{m.group(1).upper()} {m.group(2)}"
    return "(escopo global do programa)"


def extract_mod_blocks(text):
    """Encontra todos os blocos BEGIN/END OF MODIFICATION num fonte já
    remediado, e monta um resumo de cada um (onde fica, linhas, trecho de código
    com um pouco de contexto ao redor) — para a IA descrever cada alteração
    sem precisar reler o arquivo inteiro."""
    lines = text.splitlines()
    blocks = []
    i = 0
    while i < len(lines):
        if BEGIN_MOD.match(lines[i]):
            start = i
            j = i + 1
            while j < len(lines) and not END_MOD.match(lines[j]):  # procura o END correspondente
                j += 1
            end = j if j < len(lines) else len(lines) - 1
            loc = find_location(lines, start)
            # pega um pouco de contexto antes/depois do bloco (não só o bloco cru)
            ctx_start = max(0, start - CONTEXT_LINES)
            ctx_end = min(len(lines), end + CONTEXT_LINES + 1)
            snippet = "\n".join(lines[ctx_start:ctx_end])
            if len(snippet) > 4000:               # bloco gigante => corta e avisa
                snippet = snippet[:4000] + "\n... (bloco truncado) ..."
            blocks.append({"location": loc, "start_line": start + 1,
                           "end_line": end + 1, "snippet": snippet})
            i = end + 1                            # pula para depois deste bloco
        else:
            i += 1
    return blocks


def extract_signature(text, tipo):
    """Monta uma 'assinatura' resumida do objeto: só as linhas de declaração
    (REPORT, CLASS, METHODS, IMPORTING/EXPORTING, etc.), sem o corpo da lógica.
    Serve para dar à IA uma visão da estrutura do objeto sem mandar o código
    inteiro (mais barato). Limita a 120 linhas para não crescer demais."""
    picked = []
    keep = re.compile(
        r'^\s*(REPORT|PROGRAM|CLASS|INTERFACE|METHODS?|FUNCTION|FORM|MODULE|'
        r'DEFINE\s+VIEW|DEFINE\s+ROOT\s+VIEW|DEFINE\s+BEHAVIOR|MANAGED|UNMANAGED|'
        r'PROJECTION|DEFINE\s+SERVICE|EXPOSE|SELECT\s+FROM|ASSOCIATION|'
        r'PUBLIC\s+SECTION|PROTECTED\s+SECTION|IMPORTING|EXPORTING|RETURNING|'
        r'CHANGING|RAISING|TYPES|DATA|CONSTANTS)\b', re.IGNORECASE)
    for ln in text.splitlines():
        if keep.match(ln):
            picked.append(ln.rstrip())
        if len(picked) >= 120:
            picked.append("... (assinatura truncada) ...")
            break
    return "\n".join(picked)


def detect_dependencies(text):
    """Lista outros objetos SAP que este código parece usar: tabelas/CDS lidas
    em SELECT ... FROM, INCLUDEs, classes instanciadas (CREATE OBJECT/NEW/TYPE
    REF TO) e Function Modules chamados. Um `set` (conjunto) é usado para
    juntar tudo sem repetir nomes; no fim devolve uma lista ordenada, sem
    "ruído" (palavras genéricas de tipo) e limitada a 40 itens."""
    deps = set()
    for m in re.finditer(r'\bFROM\s+([a-zA-Z_/]\w+)', text, re.IGNORECASE):
        deps.add(m.group(1).upper())
    for m in re.finditer(r'\bINCLUDE\s+([a-zA-Z_/]\w+)', text, re.IGNORECASE):
        deps.add(m.group(1).upper())
    for m in re.finditer(r'(?:CREATE\s+OBJECT|TYPE\s+REF\s+TO|NEW)\s+([a-zA-Z_/]\w+)', text, re.IGNORECASE):
        deps.add(m.group(1).upper())
    for m in re.finditer(r'\bCALL\s+FUNCTION\s+[\'"]([\w/]+)[\'"]', text, re.IGNORECASE):
        deps.add(m.group(1).upper())
    # palavras genéricas de tipo/coleção que não são "dependências" de verdade
    noise = {"REF", "TABLE", "STANDARD", "SORTED", "HASHED", "DATA", "STRING"}
    return sorted(d for d in deps if d not in noise and len(d) > 2)[:40]


# ----------------------------------------------------------------------------
# Fatos determinísticos: tela de seleção, TVARV, BRF, objetos de autorização
# ----------------------------------------------------------------------------

_REPORT_TYPES = {"Report/Programa"}   # só Reports/Programas têm tela de seleção
# INCLUDE nome_do_include. (ignora includes "de sistema", tipo <ICON>)
_INCLUDE_STMT = re.compile(r"^\s*INCLUDE\s+(?:<[\w/]+>|([\w/]+))\s*\.", re.IGNORECASE | re.MULTILINE)
# O evento AT SELECTION-SCREEN, que indica que a tela de seleção tem lógica associada
_AT_SEL_SCREEN = re.compile(r"^\s*AT\s+SELECTION-SCREEN\b", re.IGNORECASE | re.MULTILINE)
# A tabela de variáveis TVARV/TVARVC, usada em ABAP para guardar parâmetros configuráveis
_TVARV_TABLE = re.compile(r"\bTVARV[C]?\b", re.IGNORECASE)
# O NOME da variável TVARV sendo lida: ou 'TVARV(C)-NAME = ...', ou um simples
# 'name = ...' / 'nome = ...' próximo à referência à tabela.
_TVARV_ASSIGN = re.compile(
    r"""TVARV[C]?\s*-\s*name\s*=\s*'([^']+)'"""
    r"""|(?:\b(?:name|nome)\b\s*(?:=|EQ)\s*'([^']+)')""",
    re.IGNORECASE,
)
# Palavras/classes que indicam uso do motor de regras BRF+ (Business Rule Framework)
_BRF_HINT = re.compile(
    r"BRF\+|BRFPLUS|BRF_PLUS|(?<![A-Z0-9_])BRF(?![A-Z0-9_+])|(?<![A-Z0-9_])BTF(?![A-Z0-9_])|"
    r"CL_FDT_\w+|IF_FDT_\w+|CL_BTF_\w+|FDT_FUNCTION|FDT_APPLICATION",
    re.IGNORECASE,
)
# AUTHORITY-CHECK OBJECT 'OBJETO_DE_AUTORIZACAO'
_AUTH_CHECK = re.compile(r"AUTHORITY-CHECK\s+OBJECT\s+'([^']+)'", re.IGNORECASE)
# Dentro de um AUTHORITY-CHECK: ID 'CAMPO_AUTH' FIELD valor
_AUTH_ID = re.compile(r"\bID\s+'([^']+)'\s+FIELD\s+(\S+)", re.IGNORECASE)
# pfcg_auth(...) — a forma como o RAP declara checagem de autorização (DCL)
_PFCG_AUTH = re.compile(r"pfcg_auth\s*\(\s*([^)]+)\)", re.IGNORECASE)
# O início de uma declaração PARAMETERS ou SELECT-OPTIONS (com ':' opcional para
# o formato "encadeado" do ABAP, ex.: 'PARAMETERS: p_a ..., p_b ... .')
_FIELD_DECL = re.compile(
    r"^\s*(?:PARAMETERS?|SELECT-OPTIONS?)\s*:?\s*",
    re.IGNORECASE,
)
# Um campo da tela: NOME TYPE|LIKE|FOR referência (ex.: 'p_data TYPE sy-datum')
_FIELD_TOKEN = re.compile(
    r"(\w+)\s+(TYPE|LIKE|FOR)\s+([\w/\-]+(?:-[\w]+)?)",
    re.IGNORECASE,
)
# SELECTION-SCREEN BEGIN OF BLOCK nome — agrupa campos num "quadro" na tela
_BLOCK_BEGIN = re.compile(
    r"SELECTION-SCREEN:?\s+BEGIN\s+OF\s+BLOCK\s+(\w+)",
    re.IGNORECASE,
)


def _is_code_comment(line):
    """True se a linha inteira é um comentário (ABAP usa '*' no início da
    linha ou '\"' em qualquer ponto) ou está vazia — ou seja, não tem código."""
    s = line.lstrip()
    return (not s) or s.startswith("*") or s.startswith('"') or s.startswith("//")


def _split_inline_comment(line):
    """Separa código ABAP do comentário `\" texto` fora de aspas simples."""
    in_squote = False
    i = 0
    while i < len(line):
        ch = line[i]
        if ch == "'":
            in_squote = not in_squote
        elif ch == '"' and not in_squote:
            return line[:i].rstrip(), line[i + 1:].strip()
        i += 1
    return line.rstrip(), ""


def _iter_code_lines(text):
    """Percorre o texto linha por linha e devolve só as linhas com código de
    verdade (número da linha, código sem comentário, comentário separado).

    É um "gerador" (usa `yield` em vez de `return`): em vez de montar uma lista
    inteira na memória, ele entrega uma linha de cada vez, sob demanda, para
    quem estiver usando um `for ... in _iter_code_lines(texto)`."""
    for i, raw in enumerate(text.splitlines(), 1):
        if _is_code_comment(raw):
            continue
        code, comment = _split_inline_comment(raw)
        if not code.strip():
            continue
        yield i, code, comment


def _find_include_file(include_name, src_path):
    """Procura no disco o arquivo-fonte de um INCLUDE citado por nome (ex.:
    'zprog_top' → 'zprog_top.abap'). Sobe a partir da pasta do arquivo atual
    até achar a pasta de códigos (DIR_CODIGOS) e procura o arquivo lá dentro
    (e em toda a árvore abaixo dela)."""
    base = (include_name or "").strip()
    if not base:
        return None
    start_dir = os.path.dirname(os.path.abspath(src_path))
    search_roots = [start_dir]
    p = start_dir
    want_cod = env_name("DIR_CODIGOS")
    # sobe até 8 níveis de pasta procurando a pasta "02_CODIGOS" (ou o nome customizado)
    for _ in range(8):
        if os.path.basename(p).lower() == want_cod.lower():
            if p not in search_roots:
                search_roots.append(p)
            break
        parent = os.path.dirname(p)
        if parent == p:      # chegou na raiz do sistema de arquivos => para
            break
        p = parent
    want = {base.lower() + ext for ext in (".abap", ".txt", ".prog", ".src", ".asinc", ".asprog")}
    for root in search_roots:
        for dirpath, _, files in os.walk(root):   # os.walk: percorre a pasta e todas as subpastas
            for fn in files:
                if fn.lower() in want:
                    return os.path.join(dirpath, fn)
    return None


def read_included_sources(text, src_path):
    """Lê includes locais de um Report (um nível) — a tela costuma estar no TOP."""
    if not src_path:
        return []
    out = []
    seen = set()
    for m in _INCLUDE_STMT.finditer(text):
        name = m.group(1)
        if not name:
            continue
        path = _find_include_file(name, src_path)
        if not path or path in seen:
            continue
        seen.add(path)
        body = read_text(path)
        if body:
            out.append((os.path.basename(path), body))
    return out


def extract_selection_fields(text):
    """PARAMETERS / SELECT-OPTIONS / blocos da tela de seleção.

    O ABAP permite declarar um campo numa linha ('PARAMETERS p_a TYPE i.') OU
    encadear vários numa declaração só, terminando em ponto só na última linha
    ('PARAMETERS: p_a TYPE i, p_b TYPE c.'). Por isso o laço usa duas variáveis
    de "estado" (`kind` = o tipo de campo atual; `pending` = ainda não chegou
    o ponto final) para saber se a linha atual continua uma declaração aberta."""
    fields = []
    seen = set()
    kind = None
    pending = False
    for _i, code, comment in _iter_code_lines(text):
        bm = _BLOCK_BEGIN.search(code)
        if bm:                                   # achou um BEGIN OF BLOCK => registra o bloco
            bid = bm.group(1)
            if bid.lower() not in seen:
                seen.add(bid.lower())
                fields.append({"kind": "BLOCK", "name": bid, "ref": "", "comment": comment})
        stripped = code.strip()
        if _AT_SEL_SCREEN.match(stripped):
            # saiu da área de declaração de campos e entrou em lógica de evento
            pending = False
            kind = None
            continue
        if _FIELD_DECL.match(stripped):
            # começou uma nova declaração PARAMETERS/SELECT-OPTIONS
            kind = "SELECT-OPTIONS" if re.match(r"^\s*SELECT-OPTIONS?", stripped, re.I) else "PARAMETERS"
            pending = True
            rest = _FIELD_DECL.sub("", stripped, count=1)   # tira a palavra-chave, sobra só os campos
        elif pending:
            # linha de continuação de uma declaração encadeada (formato "PARAMETERS: a, b, ...")
            rest = stripped
        else:
            continue
        for name, rel, ref in _FIELD_TOKEN.findall(rest):
            key = name.lower()
            if key in seen:
                continue
            seen.add(key)
            fields.append({
                "kind": kind or "PARAMETERS",
                "name": name,
                "ref": f"{rel.upper()} {ref}",
                "comment": comment,
            })
        pending = not stripped.endswith(".")    # sem ponto final => a declaração continua na próxima linha
        if not pending:
            kind = None
    return fields


def extract_tvarv(text):
    """Variáveis TVARV/TVARVC referenciadas no fonte."""
    hits = []
    seen = set()
    lines = text.splitlines()
    for i, raw in enumerate(lines):
        if _is_code_comment(raw):
            continue
        code, _c = _split_inline_comment(raw)
        if not _TVARV_TABLE.search(code):
            continue
        # olha uma "janela" das próximas 4 linhas, porque o nome da variável
        # TVARV costuma aparecer numa linha separada da referência à tabela
        window = "\n".join(lines[i:min(len(lines), i + 4)])
        names = []
        for groups in _TVARV_ASSIGN.findall(window):
            names.extend(n for n in groups if n)
        if not names:
            # nada capturado pelo padrão específico => tenta um padrão mais solto:
            # qualquer texto entre aspas simples que pareça um nome de variável Z
            names = re.findall(r"'([A-Z][A-Z0-9_/\-]{2,})'", window, re.IGNORECASE)
        skip = {"EQ", "NE", "GT", "LT", "GE", "LE", "IN", "LOW", "HIGH", "NAME", "TYPE", "SINGLE"}
        names = [n for n in names if n.upper() not in skip
                 and (n.upper().startswith("Z") or "_" in n)]
        for name in names:
            key = name.upper()
            if key in seen:
                continue
            seen.add(key)
            hits.append({"name": key, "line": i + 1, "table": "TVARVC" if re.search(r"TVARVC", code, re.I) else "TVARV"})
        if not names:
            key = f"?line{i+1}"
            if key not in seen:
                seen.add(key)
                hits.append({"name": "", "line": i + 1, "table": "TVARVC" if re.search(r"TVARVC", code, re.I) else "TVARV"})
    return hits


def extract_brf(text):
    """Referências a BRF/BRF+/BTF no fonte."""
    hits = []
    seen = set()
    for i, code, comment in _iter_code_lines(text):
        m = _BRF_HINT.search(code)
        if not m:
            continue
        token = m.group(0)
        # se houver algo entre aspas simples na linha, é provavelmente o nome
        # da regra/aplicação BRF; senão usa o próprio termo encontrado.
        quoted = re.findall(r"'([^']{3,})'", code)
        label = quoted[0] if quoted else token
        key = label.upper()
        if key in seen:
            continue
        seen.add(key)
        hits.append({"token": token, "name": label, "line": i, "comment": comment})
    return hits


def extract_auth_objects(text):
    """AUTHORITY-CHECK OBJECT e pfcg_auth (DCL).

    Aqui trabalhamos direto na STRING inteira do arquivo (em vez de linha a
    linha), porque um AUTHORITY-CHECK pode se espalhar por várias linhas antes
    do ponto final. finditer devolve todas as ocorrências do padrão no texto."""
    hits = []
    seen = set()
    for m in _AUTH_CHECK.finditer(text):
        start = m.start()
        # ignora se a ocorrência está em linha de comentário: acha o início da
        # linha (depois da última quebra \n antes daqui) para conferir
        line_start = text.rfind("\n", 0, start) + 1
        line = text[line_start:text.find("\n", start)]
        if _is_code_comment(line):
            continue
        # pega um pedaço do texto a partir daqui e corta no primeiro ponto
        # final (fim do comando ABAP), com um limite de segurança de 600 chars
        chunk = text[start:start + 600]
        dot = chunk.find(".")
        stmt = chunk[:dot] if dot != -1 else chunk[:250]
        obj = m.group(1).upper()
        ids = [{"id": a, "field": b.rstrip(".,")} for a, b in _AUTH_ID.findall(stmt)]
        key = ("AUTHORITY-CHECK", obj, tuple(d["id"] for d in ids))
        if key in seen:
            continue
        seen.add(key)
        hits.append({"kind": "AUTHORITY-CHECK", "object": obj, "ids": ids})
    for m in _PFCG_AUTH.finditer(text):
        line_start = text.rfind("\n", 0, m.start()) + 1
        line = text[line_start:text.find("\n", m.start())]
        if _is_code_comment(line):
            continue
        # pfcg_auth('OBJETO', campo1 = valor1, campo2 = valor2): o primeiro
        # item entre parênteses é o objeto, os seguintes são pares campo=valor
        parts = [p.strip() for p in m.group(1).split(",") if p.strip()]
        if not parts:
            continue
        obj = re.sub(r"['\"]", "", parts[0]).upper()
        ids = []
        for p in parts[1:]:
            p = p.strip().strip("'\"")
            if "=" in p:
                k, v = p.split("=", 1)
                ids.append({"id": k.strip().upper(), "field": v.strip().strip("'\"")})
            else:
                ids.append({"id": p.upper(), "field": ""})
        key = ("pfcg_auth", obj, tuple(d["id"] for d in ids))
        if key in seen:
            continue
        seen.add(key)
        hits.append({"kind": "pfcg_auth", "object": obj, "ids": ids})
    return hits


def analyze_extras(text, tipo, src_path=None):
    """Extrai tela (Report), TVARV, BRF e objetos de autorização do fonte."""
    src = text or ""
    applies = tipo in _REPORT_TYPES
    fields = []
    from_includes = []
    has_events = bool(_AT_SEL_SCREEN.search(src))
    if applies:
        fields.extend(extract_selection_fields(src))
        for inc_name, body in read_included_sources(src, src_path):
            extra = extract_selection_fields(body)
            if extra:
                from_includes.append(inc_name)
                fields.extend(extra)
        seen = set()
        uniq = []
        for f in fields:
            k = (f["kind"], f["name"].lower())
            if k in seen:
                continue
            seen.add(k)
            uniq.append(f)
        fields = uniq
    return {
        "selection_screen": {
            "applies": applies,
            "has_events": has_events,
            "fields": fields,
            "from_includes": from_includes,
        },
        "tvarv_variables": extract_tvarv(src),
        "brf_refs": extract_brf(src),
        "auth_objects": extract_auth_objects(src),
    }


def _md_pipe(value):
    """Deixa um valor seguro para entrar numa célula de tabela Markdown: escapa
    o caractere '|' (que separaria colunas por engano) e junta tudo numa linha."""
    return str(value or "").replace("|", "\\|").replace("\n", " ").strip()


def format_extra_sections(obj):
    """Markdown das SEÇÕES 4, 5, 5.3 e 6 (fatos do Python, sem LLM)."""
    ss = obj.get("selection_screen") or {}
    lines = ["### SEÇÃO 4 – Tela de Seleção"]
    if not ss.get("applies"):
        lines += ["Não se aplica (objeto não é Report/Programa).", ""]
    else:
        fields = ss.get("fields") or []
        if not fields and not ss.get("has_events"):
            lines += ["Não foi identificada tela de seleção (PARAMETERS / "
                      "SELECT-OPTIONS / SELECTION-SCREEN) neste programa.", ""]
        else:
            #lines.append("**Há tela de seleção:** sim")
            incs = ss.get("from_includes") or []
            if incs:
                lines.append("Parâmetros lidos também do(s) include(s): "
                             + ", ".join(f"`{n}`" for n in incs))
            if fields:
                lines += ["",
                          "| Tipo | Nome | Referência | Observação |",
                          "|------|------|------------|------------|"]
                for f in fields:
                    lines.append(
                        f"| {f.get('kind', '')} | `{f.get('name', '')}` | "
                        f"{_md_pipe(f.get('ref'))} | {_md_pipe(f.get('comment'))} |"
                    )
            elif ss.get("has_events"):
                lines.append("Há eventos `AT SELECTION-SCREEN`, mas os "
                             "PARAMETERS/SELECT-OPTIONS não estão neste fonte "
                             "(podem estar no include TOP).")
            lines.append("")

    lines.append("### SEÇÃO 5 – TVARV")
    tvarv = obj.get("tvarv_variables") or []
    if not tvarv:
        lines += ["Não foram encontradas referências à tabela TVARV/TVARVC.", ""]
    else:
        lines += ["Variáveis de TVARV/TVARVC identificadas no código:", "",
                  "| Tabela | Variável | Linha |",
                  "|--------|----------|-------|"]
        for h in tvarv:
            var = f"`{h['name']}`" if h.get("name") else "_(nome não identificado)_"
            lines.append(f"| {h.get('table', 'TVARV')} | {var} | {h.get('line', '')} |")
        lines.append("")

    lines.append("### SEÇÃO 5.3 – BRF")
    brf = obj.get("brf_refs") or []
    if not brf:
        lines += ["Não foram encontradas referências a BRF/BRF+/BTF.", ""]
    else:
        lines += ["Referências a BRF/BRF+/BTF identificadas no código:", "",
                  "| Token | Identificador | Linha |",
                  "|-------|---------------|-------|"]
        for h in brf:
            lines.append(
                f"| `{_md_pipe(h.get('token'))}` | `{_md_pipe(h.get('name'))}` | {h.get('line', '')} |"
            )
        lines.append("")

    lines.append("### SEÇÃO 6 – Objetos de Autorização")
    auths = obj.get("auth_objects") or []
    if not auths:
        lines += ["Não foram encontradas referências a objetos de autorização "
                  "(AUTHORITY-CHECK / pfcg_auth).", ""]
    else:
        lines += ["Objetos de autorização identificados no código:", "",
                  "| Origem | Objeto | Campos |",
                  "|--------|--------|--------|"]
        for h in auths:
            campos = ", ".join(
                f"{d['id']}" + (f"={d['field']}" if d.get("field") else "")
                for d in (h.get("ids") or [])
            ) or "—"
            lines.append(f"| {h.get('kind', '')} | `{h.get('object', '')}` | {_md_pipe(campos)} |")
        lines.append("")
    return lines


# Extensões de arquivo que consideramos "código-fonte SAP" ao varrer a pasta
# de códigos — tudo que não tiver uma dessas extensões é ignorado.
CODE_EXTS = (".abap", ".txt", ".src", ".prog", ".ddls", ".asddls", ".bdef",
             ".asbdef", ".srvd", ".asrvd", ".srvb", ".asrvb", ".ddlx",
             ".asddlx", ".dcl", ".amdp",
             ".aclass", ".asinc", ".asfunc", ".asprog")

# Compound ADT suffixes — strip the whole chain so .clas.abap → object name.
# (O Eclipse/ADT exporta arquivos com extensão "composta", ex.: 'ZCL_X.clas.abap'.
#  Aqui listamos essas combinações para tirar tudo de uma vez e sobrar só 'ZCL_X'.)
_COMPOUND_SUFFIXES = (
    ".clas.abap", ".intf.abap", ".ddls.asddls", ".srvd.asrvds", ".srvb.asrvb",
    ".bdef.asbdef", ".ddlx.asddlx", ".asddls", ".asbdef", ".asrvd", ".asrvb",
    ".asddlx", ".clas", ".intf", ".abap", ".txt", ".src", ".prog", ".ddls",
    ".bdef", ".srvd", ".srvb", ".ddlx", ".dcl", ".amdp",
    ".aclass", ".asinc", ".asfunc", ".asprog",
)


def object_stem(filename):
    """Nome do objeto sem extensão ADT composta (.clas.abap, .ddls.asddls, …).

    Testa os sufixos do MAIS longo para o MAIS curto (sorted com key=len,
    reverse=True) para não cortar errado — ex.: se testasse '.abap' antes de
    '.clas.abap', 'ZCL_X.clas.abap' viraria 'ZCL_X.clas' em vez de 'ZCL_X'."""
    name = os.path.basename(filename)
    lower = name.lower()
    for suf in sorted(_COMPOUND_SUFFIXES, key=len, reverse=True):
        if lower.endswith(suf):
            return name[:len(name) - len(suf)]
    return os.path.splitext(name)[0]


def read_text(path):
    """Lê um arquivo de texto tentando várias codificações, na ordem: UTF-8
    (o padrão moderno), depois Latin-1 e CP1252 (comuns em exportações
    Windows/SAP antigas). Devolve None se nenhuma delas funcionar."""
    for enc in ("utf-8", "latin-1", "cp1252"):
        try:
            with open(path, "r", encoding=enc) as f:
                return f.read()
        except (UnicodeDecodeError, OSError):
            continue
    return None


def gather_code_files(codigos_dir):
    """Lista, em ordem alfabética, todos os arquivos de código (extensão em
    CODE_EXTS) dentro da pasta de códigos, incluindo subpastas."""
    files = []
    for dirpath, _, filenames in os.walk(codigos_dir):
        for fn in filenames:
            if fn.startswith("~$") or fn.startswith("."):    # ignora temporários do Office e ocultos
                continue
            full = os.path.join(dirpath, fn)
            ext = os.path.splitext(fn)[1].lower()
            if ext in CODE_EXTS:
                files.append(full)
    return sorted(files)


def _ef_skip_names():
    """Nomes de arquivo que NUNCA devem ser tratados como a EF em si (são os
    arquivos auxiliares de regras/requisitos que também moram na pasta da EF)."""
    return {
        env_name("FILE_REGRAS").lower(),
        env_name("FILE_REQUISITOS").lower(),
    }


def _is_ef_skip(fn):
    """True se este arquivo NÃO deve ser considerado candidato a EF (é
    temporário, é um dos arquivos auxiliares, ou parece ser o modelo/template)."""
    low = fn.lower()
    if fn.startswith("~$") or fn.startswith("."):
        return True
    if low in _ef_skip_names():
        return True
    if "template" in low:
        return True
    return False


def find_ef_file(ef_dir):
    """Escolhe qual arquivo da pasta 01_EF é a Especificação Funcional.

    Primeiro tenta achar candidatos "normais" (ignorando template/auxiliares).
    Se não sobrar nenhum, cede um pouco e aceita até um arquivo de template
    (último recurso, melhor que não ter EF nenhuma). Entre os candidatos,
    prioriza .docx, depois .pdf, .md e .txt (a ordem típica de qualidade de
    extração de texto)."""
    if not os.path.isdir(ef_dir):
        return None, None
    candidates = [os.path.join(ef_dir, fn) for fn in os.listdir(ef_dir)
                  if os.path.isfile(os.path.join(ef_dir, fn)) and not _is_ef_skip(fn)]
    if not candidates:
        # último recurso: aceita template se for o único .docx/.pdf
        candidates = [os.path.join(ef_dir, fn) for fn in os.listdir(ef_dir)
                      if os.path.isfile(os.path.join(ef_dir, fn))
                      and not fn.startswith("~$")
                      and fn.lower() not in _ef_skip_names()]
    if not candidates:
        return None, None
    prio = {".docx": 0, ".pdf": 1, ".md": 2, ".txt": 3}
    candidates.sort(key=lambda p: (prio.get(os.path.splitext(p)[1].lower(), 9),
                                   os.path.basename(p).lower()))
    ef = candidates[0]
    return ef, os.path.splitext(ef)[1].lower().lstrip(".")


def read_docx_text(path):
    """Extrai o texto de um arquivo Word (.docx): parágrafos normais + o
    conteúdo de tabelas (cada linha de tabela vira uma linha 'célula | célula').

    Devolve sempre uma tupla (texto, erro) — se der certo, erro é None; se
    falhar, texto é None e erro tem a mensagem (esse é o padrão usado em todas
    as funções read_*_text deste arquivo, para nunca lançar exceção pro
    chamador e sempre poder registrar o motivo no relatório)."""
    try:
        import docx     # biblioteca externa 'python-docx'; só importada se for realmente usada
    except ImportError:
        return None, "python-docx não instalado (pip install python-docx)"
    try:
        doc = docx.Document(path)
        parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
        for tbl in doc.tables:
            for row in tbl.rows:
                cells = []
                for c in row.cells:
                    txt = c.text.strip()
                    # dedup de célula mesclada (python-docx repete o texto)
                    if txt and (not cells or cells[-1] != txt):
                        cells.append(txt)
                line = " | ".join(cells)
                if line:
                    parts.append(line)
        return "\n".join(parts), None
    except Exception as e:
        return None, f"falha ao ler docx: {e}"


def read_pdf_text(path):
    """Extrai o texto de um PDF. Tenta primeiro a biblioteca pdfplumber (mais
    precisa em PDFs com tabelas); se falhar ou não estiver instalada, cai para
    pypdf como alternativa. Mesmo padrão de retorno (texto, erro) das outras
    funções read_*_text."""
    try:
        import pdfplumber
        out = []
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                t = page.extract_text() or ""
                if t.strip():
                    out.append(t)
        if out:
            return "\n".join(out), None
    except Exception:
        pass    # pdfplumber falhou (ou não achou texto) => tenta o plano B abaixo
    try:
        import pypdf
        reader = pypdf.PdfReader(path)
        out = [(pg.extract_text() or "") for pg in reader.pages]
        return "\n".join(t for t in out if t.strip()), None
    except Exception as e:
        return None, f"falha ao ler pdf: {e}"


def read_ef_text(path, ext):
    """Escolhe o leitor certo (texto simples, .docx ou .pdf) de acordo com a
    extensão do arquivo da EF, e devolve (texto, erro)."""
    if ext in ("md", "txt"):
        return read_text(path), None
    if ext == "docx":
        return read_docx_text(path)
    if ext == "pdf":
        return read_pdf_text(path)
    txt = read_text(path)
    if txt is not None:
        return txt, None
    return None, f"tipo de EF não suportado para extração automática: .{ext}"


def _clean_field(field):
    """Limpa o nome de um campo extraído das regras: tira '**negrito**'
    markdown, tira crases de código, e se o texto tiver 'Rótulo: alguma coisa',
    fica só com o 'Rótulo' (remove o placeholder de exemplo, ex.: 'Projeto:
    EF.Projeto' vira só 'Projeto')."""
    field = re.sub(r'^\*\*|\*\*$', '', field).strip()
    field = field.strip("`").strip()
    # remove placeholder do tipo "Projeto: EF.Projeto" -> "Projeto"
    if ":" in field:
        left = field.split(":", 1)[0].strip()
        if 1 < len(left) < 60:
            field = left
    return field.strip()


# Campos do cabeçalho padrão da EF (tabela "Rótulo: valor" da 1ª página),
# usados quando não existe regras_extrair.md na pasta da EF.
DEFAULT_RULE_FIELDS = [
    "Projeto", "Fase do Projeto", "Autor", "Módulo", "Cenário empresarial",
    "Processo", "ID GAP", "Descrição GAP", "Transações",
]


def parse_rule_fields(rules_text):
    """Extrai a lista de campos pedidos.

    Estratégia:
      1) Se houver itens de lista (bullets '- '/'* '/'+ ' ou numerados),
         usa APENAS esses itens como campos (formato recomendado, sem ambiguidade).
      2) Caso contrário, faz fallback tolerante: separa pares por TAB e reconhece
         'Rótulo: valor' e o placeholder 'Rótulo: EF.Rótulo'. Ignora linhas de
         instrução (sem ':' + valor) e cabeçalhos markdown.
    """
    if not rules_text:
        return []

    # ---- 1) bullets/numerados têm prioridade ----
    bullets = []
    for raw in rules_text.splitlines():
        line = raw.strip()
        if not line or line.startswith("#"):
            continue
        m = re.match(r'^(?:[-*+]\s+|\d+[.)]\s+)(.*)$', line)
        if m:
            f = _clean_field(m.group(1))
            if f and f.lower() not in {b.lower() for b in bullets}:
                bullets.append(f)
    if bullets:
        return bullets

    # ---- 2) fallback: pares separados por TAB + placeholder EF. ----
    fields = []
    for raw in rules_text.splitlines():
        line = raw.rstrip()
        stripped = line.strip()
        if not stripped or stripped.startswith("#"):
            continue
        # ignora linhas de instrução conhecidas
        low = stripped.lower()
        if low.startswith(("extrair da tabela", "deve coletar", "extrair da ef",
                           "extrair os campos", "cada campo", "fonte")):
            continue
        # separa dois (ou mais) campos por TAB, ou por 2+ espaços
        segments = re.split(r'\t+|\s{2,}', line)
        for seg in segments:
            seg = seg.strip()
            # 'Rótulo:' (mesmo vazio) ou 'Rótulo: valor'
            m = re.match(r'^([^:]{2,60}):\s*(.*)$', seg)
            if not m:
                continue
            field = _clean_field(seg)
            if field and field.lower() not in {f.lower() for f in fields}:
                fields.append(field)
    return fields


def extract_ef_values(ef_text, fields):
    """Extrai deterministicamente os valores 'Rótulo: Valor' do texto da EF.

    Para cada campo, procura 'Rótulo:' no texto e captura o valor até o próximo
    rótulo conhecido ou até o fim da célula (delimitada por '|' ou quebra de linha).
    Retorna dict {campo: valor}. Campos não encontrados ficam de fora.
    """
    if not ef_text or not fields:
        return {}
    # rótulos ordenados do mais longo para o mais curto evita que 'GAP' capture
    # dentro de 'ID GAP', ou 'Projeto' dentro de 'Fase do Projeto'
    labels_sorted = sorted(fields, key=len, reverse=True)
    alt = "|".join(re.escape(f) for f in labels_sorted)
    # só no início de linha ou de célula de tabela — evita "mais de um módulo:"
    label_re = re.compile(r'(?:^|\n|\|)\s*(' + alt + r')\s*:\s*', re.IGNORECASE)
    matches = list(label_re.finditer(ef_text))
    values = {}
    for i, m in enumerate(matches):
        label = m.group(1)
        start = m.end()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(ef_text)
        raw = ef_text[start:end]
        # o valor não cruza fronteira de célula/linha
        raw = re.split(r'[|\n]', raw)[0].strip()
        canon = next((f for f in fields if f.lower() == label.lower()), label)
        if canon not in values and raw:
            values[canon] = raw
    return values


def build_requisitos(fields, values, date_display, ef_rel, ef_error):
    """Monta o Markdown do arquivo de requisitos (requisitos_extraidos.md):
    um título por campo pedido, com o valor já extraído da EF, ou o
    placeholder `_(a preencher)_` quando o Python não conseguiu achar."""
    found = sum(1 for f in fields if f in values)
    lines = [
        "# Requisitos Extraídos da EF", "",
        f"Data de geração: {date_display}",
        f"Fonte (EF): {ef_rel or 'NÃO ENCONTRADA'}",
        f"Campos: {len(fields)} | Preenchidos automaticamente: {found}", "",
        "> Valores extraídos automaticamente da tabela da EF. "
        "Campos marcados `_(a preencher)_` não foram localizados — "
        "verifique na EF e preencha, ou registre `Não informado na EF`.", "",
    ]
    if ef_error:
        lines.append(f"> ATENÇÃO: não foi possível ler a EF automaticamente ({ef_error}). "
                     "Leia a EF com a skill apropriada (docx/pdf) e preencha manualmente.")
        lines.append("")
    if not fields:
        lines.append("_(Nenhum campo reconhecido em regras_extrair.md — "
                     "verifique o formato do arquivo de regras.)_")
    else:
        for fld in fields:
            lines.append(f"## {fld}")
            lines.append(values.get(fld, "_(a preencher)_"))
            lines.append("")
    return "\n".join(lines)


def make_doc_name(base, date_str, used):
    """Gera o nome do arquivo de documentação de um objeto, ex.:
    'ZCL_MEU_OBJETO_19082026_DOC'. Se esse nome já tiver sido usado (dois
    arquivos com o mesmo nome sem extensão), acrescenta _2, _3, ... até achar
    um nome livre — o parâmetro `used` é o conjunto de nomes já escolhidos."""
    xxx = object_stem(base).upper()
    xxx = re.sub(r'[^A-Z0-9_]', '_', xxx)   # troca qualquer caractere "estranho" por _
    candidate = f"{xxx}_{date_str}_DOC"
    n = 2
    while candidate in used:
        candidate = f"{xxx}_{date_str}_DOC_{n}"
        n += 1
    used.add(candidate)
    return candidate


def skeleton(obj, date_display, gap_ef):
    """Monta o Markdown 'esqueleto' do documento técnico de UM objeto: os
    dados que o Python já sabe (nome, tipo, status) mais os títulos de seção
    prontos com placeholders `_(a preencher)_` — é esse arquivo que a IA, na
    etapa seguinte, vai abrir e completar com a prosa (o texto explicativo)."""
    status = "Remediação" if obj["remediation"] else "Objeto Novo"
    lines = [
        f"# Documentação Técnica — {obj['name_no_ext']}", "",
        f"- **Arquivo:** {obj['file_name']}",
        f"- **Tipo do objeto:** {obj['type']}",
        f"- **Classificação:** {status}",
        f"- **Data de geração:** {date_display}",
        f"- **GAP/EF relacionada:** {gap_ef}", "",
        "## 📌 RESUMO EXECUTIVO", "",
        "**Entendimento:** _(a preencher: o que o objeto faz e seu papel no GAP)_", "",
        "**Contexto de Negócio:** _(a preencher: finalidade e regra de negócio)_", "",
        "**Avaliação Geral:** _(a preencher: maturidade, pontos de atenção e lacunas)_", "",
        "## 📊 RELATÓRIO ESTRUTURADO DE VALIDAÇÃO", "",
        "### SEÇÃO 1 – Resumo do Desenvolvimento",
        "_(a preencher: tabela Status | Ponto | Observação)_", "",
        "### SEÇÃO 2 – Detalhamento do Desenvolvimento",
    ]
    if obj["remediation"]:
        lines.append(f"_(a preencher: descrever cada um dos {len(obj['mod_blocks'])} "
                     "bloco(s) BEGIN/END OF MODIFICATION — antes/depois e justificativa técnica)_")
        lines.append("")
        for k, b in enumerate(obj["mod_blocks"], 1):
            lines.append(f"#### Bloco {k} — alteração em `{b['location']}` "
                         f"(linhas {b['start_line']}–{b['end_line']})")
            lines.append("```abap")
            lines.append(b["snippet"])
            lines.append("```")
            lines.append("")
    else:
        lines.append("_(a preencher: detalhar a implementação — rotinas/métodos "
                     "principais e a regra de negócio aplicada)_")
        lines.append("")
    lines += ["### SEÇÃO 3 – Observações Técnicas",
              "_(a preencher: tabela Status | Ponto | Observação — impactos S/4HANA, "
              "Clean Core, desempenho, pontos de atenção)_", ""]
    lines += format_extra_sections(obj)
    return "\n".join(lines)


def build_objetos_md(objects, date_display):
    """Monta o Markdown de 'objetos_extraidos.md': uma tabela-resumo com todos
    os objetos encontrados, na ordem de desenvolvimento (ver TYPE_ORDER)."""
    lines = [
        "# Objetos Extraídos", "",
        f"Data de geração: {date_display}",
        f"Total de objetos: {len(objects)}",
        f"Novos: {sum(1 for o in objects if not o['remediation'])} | "
        f"Remediações: {sum(1 for o in objects if o['remediation'])}", "",
        "Objetos ordenados por tipo, na ordem correta de desenvolvimento.", "",
        "| # | Nome do Objeto | Tipo | Status | Documento gerado |",
        "|---|----------------|------|--------|------------------|",
    ]
    for i, o in enumerate(objects, 1):
        status = "Remediação" if o["remediation"] else "Novo"
        lines.append(f"| {i} | `{o['name_no_ext']}` | {o['type']} | {status} | `{o['doc_name']}.md` |")
    return "\n".join(lines)


def main():
    """O PONTO DE ENTRADA do script: orquestra toda a extração, do início ao
    fim. Em resumo, os passos são:
      1) ler os parâmetros da linha de comando e o .env;
      2) descobrir e ler a EF (Especificação Funcional);
      3) extrair os valores dos campos pedidos em regras_extrair.md;
      4) gerar requisitos_extraidos.md;
      5) classificar cada arquivo de código e montar um "esqueleto" .md por objeto;
      6) montar o _payload.json com tudo, para a IA usar na etapa seguinte;
      7) imprimir um resumo no terminal.
    Cada passo é comentado abaixo, na ordem em que acontece."""
    # argparse: lê os argumentos passados na linha de comando (ex.: --date 19082026)
    ap = argparse.ArgumentParser(description="Extrator mecânico EF/códigos SAP (sem LLM).")
    ap.add_argument("--root", default=None,
                    help="Pasta inbound alternativa (padrão: input/). Não use a pasta do extract.py.")
    ap.add_argument("--date", default=None, help="Data DDMMYYYY (default: hoje)")
    ap.add_argument("--ef-chars", type=int, default=20000,
                    help="Máx. de caracteres da EF no payload (default 20000)")
    ap.add_argument("--src-chars", type=int, default=8000,
                    help="Máx. de caracteres de código por objeto novo no payload (default 8000)")
    args = ap.parse_args()

    load_project_env()
    try:
        lay = parse_location(root=args.root)
    except LayoutError as e:
        print(f"ERRO: {e}", file=sys.stderr)
        sys.exit(2)

    root = lay.inbound
    ef_dir, ef_used = lay.ef, lay.ef_name
    cod_dir, cod_used = lay.codes, lay.codes_name
    out_dir, out_used = lay.artifacts, lay.artifacts_rel
    inbound_used = env_name("ET_INBOUND_DIR")
    outbound_used = env_name("ET_OUTBOUND_DIR")
    et_used = env_name("DIR_ET_GERADA")
    review_used = lay.review_name
    req_file = env_name("FILE_REQUISITOS")
    regras_file = env_name("FILE_REGRAS")

    if _is_inside(out_dir, lay.inbound):
        print("ERRO: recusando escrita dentro de input. "
              "A raiz do projeto é a pasta do extract.py; a saída é output/.",
              file=sys.stderr)
        sys.exit(2)

    if not os.path.isdir(cod_dir):
        print(f"ERRO: diretório não encontrado: {cod_dir} "
              f"(DIR_CODIGOS={cod_used})", file=sys.stderr)
        sys.exit(2)
    os.makedirs(out_dir, exist_ok=True)

    if args.date:
        # --date foi passado: usa exatamente o que veio (ex.: '19082026'), e
        # tenta convertê-lo para o formato bonito 'dd/mm/aaaa' só para exibição.
        date_str = args.date
        try:
            date_display = datetime.strptime(date_str, "%d%m%Y").strftime("%d/%m/%Y")
        except ValueError:
            date_display = date_str
    else:
        # sem --date: usa a data/hora de hoje
        now = datetime.now()
        date_str = now.strftime("%d%m%Y")
        date_display = now.strftime("%d/%m/%Y")

    ef_file, ef_type = find_ef_file(ef_dir)
    ef_rel = os.path.relpath(ef_file, root) if ef_file else None
    rules_path = os.path.join(ef_dir, regras_file)
    rules_text = read_text(rules_path) if os.path.isfile(rules_path) else None
    rule_fields = parse_rule_fields(rules_text) if rules_text else list(DEFAULT_RULE_FIELDS)

    ef_text, ef_error = (None, None)
    if ef_file:
        ef_text, ef_error = read_ef_text(ef_file, ef_type)
    else:
        ef_error = f"nenhum arquivo de EF encontrado em {ef_used}/"

    ef_values = extract_ef_values(ef_text, rule_fields) if ef_text else {}

    # GAP/EF relacionada — determinístico: "ID GAP / Nome da EF"
    # (função aninhada: definida aqui dentro porque só é usada nas 2 linhas
    #  seguintes; assim não polui o resto do arquivo com um nome de função extra)
    def _find_val(keys):
        """Procura, entre os valores já extraídos da EF, o primeiro cuja
        chave (nome do campo) contenha algum dos termos em `keys`."""
        for k, v in ef_values.items():
            if any(t in k.lower() for t in keys):
                return v
        return None
    id_gap = _find_val(["id gap", "gap"])
    ef_name = os.path.splitext(os.path.basename(ef_file))[0] if ef_file else None
    if id_gap and ef_name:
        gap_ef = f"{id_gap} / {ef_name}"
    elif id_gap:
        gap_ef = id_gap
    elif ef_name:
        gap_ef = ef_name
    else:
        gap_ef = f"_(preencher a partir de {req_file})_"

    requisitos_rel = None
    req_path = os.path.join(out_dir, req_file)
    # 'with open(...) as f' abre o arquivo e garante que ele é fechado
    # automaticamente ao sair do bloco, mesmo se der erro no meio.
    with open(req_path, "w", encoding="utf-8") as f:
        f.write(build_requisitos(rule_fields, ef_values, date_display, ef_rel, ef_error))
    requisitos_rel = os.path.relpath(req_path, project_dir())

    files = gather_code_files(cod_dir)
    if not files:
        print(f"ERRO: nenhum arquivo de código em {cod_dir}", file=sys.stderr)
        sys.exit(3)

    # Passo principal: para cada arquivo de código, lê, classifica e extrai
    # tudo que é determinístico (tipo, remediação, dependências, tela, etc.)
    objects = []
    used_doc_names = set()
    for path in files:
        text = read_text(path)
        if text is None:
            continue
        file_name = os.path.basename(path)
        name_no_ext = object_stem(file_name)
        tipo = classify(path, text)
        remed = is_remediation(text)
        extras = analyze_extras(text, tipo, path)
        obj = {
            "file_name": file_name, "name_no_ext": name_no_ext,
            "rel_path": os.path.relpath(path, root), "type": tipo,
            "type_rank": TYPE_ORDER_MAP.get(tipo, 120), "remediation": remed,
            "mod_blocks": extract_mod_blocks(text) if remed else [],
            "dependencies": detect_dependencies(text),
            "line_count": len(text.splitlines()),
        }
        obj.update(extras)
        obj["doc_name"] = make_doc_name(file_name, date_str, used_doc_names)
        objects.append(obj)

    # ordena pela "ordem de desenvolvimento" (TYPE_ORDER) e, dentro do mesmo
    # tipo, por nome — lambda: uma função "descartável", escrita numa linha só,
    # que diz ao sort() qual é a "chave" de ordenação de cada objeto.
    objects.sort(key=lambda o: (o["type_rank"], o["name_no_ext"].lower()))

    with open(os.path.join(out_dir, "objetos_extraidos.md"), "w", encoding="utf-8") as f:
        f.write(build_objetos_md(objects, date_display))

    # trunca o texto da EF se for grande demais, para não estourar o
    # orçamento de tokens quando a IA ler o payload
    ef_payload_text, ef_truncated = None, False
    if ef_text:
        if len(ef_text) > args.ef_chars:
            ef_payload_text = ef_text[:args.ef_chars] + "\n... (EF truncada) ..."
            ef_truncated = True
        else:
            ef_payload_text = ef_text

    payload = {
        "generated_at": date_display, "date_str": date_str,
        "ef_file": ef_rel, "ef_type": ef_type, "ef_read_error": ef_error,
        "ef_truncated": ef_truncated, "ef_text": ef_payload_text,
        "regras_extrair_file": os.path.relpath(rules_path, root) if rules_text else None,
        "regras_extrair_text": rules_text, "rule_fields": rule_fields,
        "ef_values": ef_values, "gap_ef": gap_ef,
        "et_template": resolve_et_template(),
        "requisitos_file": requisitos_rel,
        "artifacts_dir": out_used,
        "inbound_dir": inbound_used,
        "outbound_dir": outbound_used,
        "id_gap": id_gap,
        "objects": [],
    }

    # limite de contexto de código por objeto novo (controle de custo de LLM)
    src_limit = args.src_chars

    # Para cada objeto: grava o esqueleto .md e monta a "entrada" dele no
    # payload que a IA vai ler (remediação => foco nos blocos alterados;
    # objeto novo => manda a assinatura + um trecho do código-fonte).
    for obj in objects:
        with open(os.path.join(out_dir, obj["doc_name"] + ".md"), "w", encoding="utf-8") as f:
            f.write(skeleton(obj, date_display, gap_ef))
        text = read_text(os.path.join(root, obj["rel_path"]))
        entry = {
            "file_name": obj["file_name"], "name_no_ext": obj["name_no_ext"],
            "rel_path": obj["rel_path"], "type": obj["type"],
            "status": "Remediação" if obj["remediation"] else "Novo",
            "doc_file": obj["doc_name"] + ".md", "dependencies": obj["dependencies"],
            "line_count": obj["line_count"],
            "selection_screen": obj.get("selection_screen"),
            "tvarv_variables": obj.get("tvarv_variables") or [],
            "brf_refs": obj.get("brf_refs") or [],
            "auth_objects": obj.get("auth_objects") or [],
        }
        if obj["remediation"]:
            # foco nos blocos alterados — barato e suficiente
            entry["mod_blocks"] = obj["mod_blocks"]
            entry["truncated"] = False
        else:
            # objeto novo: precisa de contexto de código para descrever a regra de negócio.
            # envia o fonte inteiro se pequeno; senão trunca e sinaliza para ler o arquivo.
            entry["signature"] = extract_signature(text, obj["type"])
            if text and len(text) > src_limit:
                entry["source_excerpt"] = text[:src_limit] + "\n... (fonte truncado) ..."
                entry["truncated"] = True
            else:
                entry["source_excerpt"] = text or ""
                entry["truncated"] = False
        payload["objects"].append(entry)

    # grava o payload completo em JSON — é o arquivo que a etapa seguinte
    # (com IA) vai ler para escrever a prosa dos documentos.
    with open(os.path.join(out_dir, "_payload.json"), "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)

    # resumo final impresso no terminal, para conferência rápida
    print("=" * 60)
    print("EXTRAÇÃO CONCLUÍDA (sem LLM)")
    print("=" * 60)
    print(f"Projeto         : {project_dir()}")
    print(f"Inbound         : {root}  [{inbound_used}/]")
    print(f"  EF            : {ef_dir}")
    print(f"  Códigos       : {cod_dir}")
    print(f"Outbound        : {out_dir}")
    print(f"Pastas          : EF={ef_used}  códigos={cod_used}  "
          f"ET={et_used}  artefatos={out_used}  review={review_used}")
    print(f"Data            : {date_display} ({date_str})")
    print(f"EF detectada    : {ef_rel or 'NÃO ENCONTRADA'}")
    if ef_error:
        print(f"EF leitura      : ERRO -> {ef_error}")
    elif ef_text:
        print(f"EF leitura      : ok ({len(ef_text)} chars{', truncada no payload' if ef_truncated else ''})")
    print(f"{regras_file:<16}: {'ok' if rules_text else 'NÃO ENCONTRADO (usando campos padrão do cabeçalho da EF)'} "
          f"({len(rule_fields)} campo(s) reconhecido(s))")
    print(f"Valores da EF   : {len(ef_values)}/{len(rule_fields)} preenchidos automaticamente")
    print(f"{req_file} : gerado em {out_used}")
    print(f"Objetos         : {len(objects)} "
          f"(Novos: {sum(1 for o in objects if not o['remediation'])}, "
          f"Remediações: {sum(1 for o in objects if o['remediation'])})")
    print(f"Gerados em      : {out_dir}")
    print(f"Template ET     : {describe_et_template(payload.get('et_template'))}")
    for o in objects:
        print(f"  - {o['doc_name']}.md  [{o['type']}]")
    print("=" * 60)
    print("SAÍDAS:")
    print(f"  {out_used}/{req_file}")
    print(f"  {out_used}/objetos_extraidos.md + *_DOC.md + _payload.json")
    print(f"  {review_used}/{env_name('FILE_ACHADOS')}  (Code Inspector: lc_CapCodeReview.py)")
    print(f"  {review_used}/{env_name('FILE_PENDENCIAS')}")
    print(f"  {review_used}/{env_name('FILE_RESULTADO')}")
    print("PRÓXIMO PASSO (LLM):")
    print(f"  Passo 1: preencher {out_used}/{req_file} a partir de payload.ef_text")
    print("  Passo 3: preencher a prosa de cada *_DOC.md (análise de código)")
    print("           - via agente (Claude), OU")
    print("           - headless: python generate_docs.py")


if __name__ == "__main__":
    main()